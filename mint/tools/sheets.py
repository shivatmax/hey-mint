"""Documents into a spreadsheet: "put every invoice in Downloads into a spreadsheet",
"make a table of these receipts: date, shop, total", "list the candidates' CVs with
email, years of experience and current company".

make_spreadsheet gathers the files (a folder, a pattern, or a list), reads each one -
the text of PDFs, Word, Pages-exported and text files on the Mac; scans, photos and
screenshots go to Gemini as they are - and Gemini Flash Lite pulls the same fields
out of every file, several at a time. When the user names no columns, they are chosen
from the first few files. The result is a real .xlsx (numbers as numbers, dates as
dates, a bold frozen header, a column with the source file) in the Spreadsheets folder
of Mint's storage, and it is opened.
"""

from __future__ import annotations

import datetime as dt
import fnmatch
import logging
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

log = logging.getLogger("mint.tools.sheets")

READABLE = {".pdf", ".txt", ".md", ".csv", ".json", ".html", ".htm", ".rtf", ".doc", ".docx", ".odt", ".eml",
            ".png", ".jpg", ".jpeg", ".heic", ".webp", ".gif", ".tiff"}
IMAGES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
          ".heic": "image/heic", ".gif": "image/gif", ".tiff": "image/tiff"}
MAX_FILES = 200
MAX_BYTES = 18_000_000          # Gemini's inline limit is 20 MB per request


def _files(source: str, pattern: str = "") -> list[Path]:
    """A folder (optionally with a pattern), a single file, or several paths separated by ';' or newlines."""
    from mint.tools import harness as harness_tools
    out: list[Path] = []
    for part in re.split(r"[;\n]+", str(source or "")):
        part = part.strip()
        if not part:
            continue
        path = Path(os.path.expanduser(part))
        if not path.is_absolute():
            try:
                found, _why = harness_tools._resolve(part, must_exist=True)       # "downloads", "desktop/receipts"
            except Exception:
                found = None
            path = found or Path.home() / part
        if path.is_dir():
            for entry in sorted(path.iterdir(), key=lambda p: p.name.lower()):
                if entry.is_file() and not entry.name.startswith(".") and entry.suffix.lower() in READABLE \
                        and (not pattern or fnmatch.fnmatch(entry.name.lower(), pattern.lower())):
                    out.append(entry)
        elif path.is_file():
            out.append(path)
    return out[:MAX_FILES]


def _content(path: Path):
    """What to show the model for one file: text, or the file itself (scans, photos)."""
    from google.genai import types
    from mint.tools import harness as harness_tools
    suffix = path.suffix.lower()
    if suffix in IMAGES and path.stat().st_size < MAX_BYTES:
        return types.Part.from_bytes(data=path.read_bytes(), mime_type=IMAGES[suffix])
    text, _note = harness_tools._text_of(path)
    if text.strip():
        return text[:40_000]
    if suffix == ".pdf" and path.stat().st_size < MAX_BYTES:          # a scan: Gemini reads the pages
        return types.Part.from_bytes(data=path.read_bytes(), mime_type="application/pdf")
    return None


def _choose_columns(what: str, samples: list[tuple[Path, object]]) -> list[str]:
    from mint.core import llm
    parts: list = [f"The user wants a spreadsheet of these files: {what or 'one row per file'}. Here are a few of "
                   "them. Choose 4 to 10 column names (short, Title Case) that capture what matters and appear in "
                   "most of them - for invoices e.g. Vendor, Invoice Number, Date, Due Date, Total, Currency. "
                   'Return JSON: {"columns": [...]}']
    for path, content in samples[:3]:
        parts += [f"\n--- {path.name} ---", content if not isinstance(content, str) else content[:6000]]
    text, _ = llm.generate(parts, json_mode=True)
    columns = llm.parse_json(text).get("columns") or []
    return [str(c).strip() for c in columns if str(c).strip()][:12] or ["Title", "Date", "Summary"]


def _extract(path: Path, content, columns: list[str], what: str) -> dict:
    from mint.core import llm
    prompt = (f"Extract these fields from the document '{path.name}' ({what or 'one row per file'}): "
              f"{', '.join(columns)}. Return JSON with exactly those keys. Numbers as plain numbers (no currency "
              "signs or thousands separators), dates as YYYY-MM-DD, text short. Unknown -> null. If the file "
              "clearly holds several items that each deserve a row (a statement with many transactions), return "
              '{"rows": [ {...}, ... ]} instead.')
    try:
        text, _ = llm.generate([prompt, content], json_mode=True)
        data = llm.parse_json(text)
    except Exception as error:
        log.info("%s: %s", path.name, error)
        return {"rows": [{"_error": str(error)[:120]}]}
    rows = data.get("rows") if isinstance(data, dict) and isinstance(data.get("rows"), list) else [data]
    return {"rows": [r for r in rows if isinstance(r, dict)]}


def _cell(value):
    """Numbers as numbers and dates as dates, so the sheet can sum and sort."""
    if isinstance(value, (int, float)) or value is None:
        return value
    text = str(value).strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", text.replace(",", "")) and len(text) < 16 and not text.startswith("0"):
        try:
            return float(text.replace(",", "")) if "." in text else int(text.replace(",", ""))
        except ValueError:
            pass
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        try:
            return dt.date.fromisoformat(text)
        except ValueError:
            pass
    return text


def _write(path: Path, columns: list[str], rows: list[tuple[str, dict]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    header = [*columns, "Source file"]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E6F4F1")
    for source, row in rows:
        values = [_cell(row.get(c)) for c in columns] + [source]
        if row.get("_error"):
            values[0] = f"(could not read: {row['_error']})"
        sheet.append(values)
    for i, column in enumerate(sheet.columns, 1):
        width = max(len(str(c.value or "")) for c in column)
        sheet.column_dimensions[column[0].column_letter].width = min(60, max(10, width + 2))
        for c in column[1:]:
            if isinstance(c.value, dt.date):
                c.number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)


def make_spreadsheet(source: str, what: str = "", columns=None, pattern: str = "", name: str = "",
                     open_after: bool = True, save_to: str = "") -> str:
    from mint.core import config
    files = _files(source, pattern)
    if not files:
        return f"FAILED: found no readable files in '{source}'" + (f" matching {pattern}" if pattern else "") + "."
    started = time.monotonic()
    with ThreadPoolExecutor(6) as pool:
        contents = list(pool.map(lambda p: (p, _content(p)), files))
    usable = [(p, c) for p, c in contents if c is not None]
    skipped = [p.name for p, c in contents if c is None]
    if not usable:
        return "FAILED: none of those files had anything readable in them."
    cols = [str(c).strip() for c in (columns or []) if str(c).strip()] or _choose_columns(what, usable)
    with ThreadPoolExecutor(6) as pool:
        results = list(pool.map(lambda pc: (pc[0], _extract(pc[0], pc[1], cols, what)), usable))
    rows = [(p.name, r) for p, res in results for r in res["rows"]]
    title = re.sub(r"[^\w .-]", "", name or what or "Spreadsheet").strip()[:60] or "Spreadsheet"
    from mint.tools import saveto
    title = re.sub(r"\.(xlsx|xls|csv)$", "", title, flags=re.I).strip() or "Spreadsheet"
    out, moved = saveto.destination(config.storage("Spreadsheets") / f"{title}.xlsx", ".xlsx", save_to)
    if out.parent == config.storage("Spreadsheets"):
        out = config.storage("Spreadsheets") / f"{title} {time.strftime('%Y-%m-%d %H.%M')}.xlsx"
    _write(out, cols, rows)
    if open_after:
        import subprocess
        subprocess.run(["open", str(out)], check=False)
    failed = sum(1 for _, r in rows if r.get("_error"))
    preview = "; ".join(", ".join(f"{c}: {r.get(c)}" for c in cols[:4]) for _, r in rows[:3])
    return (f"Made a spreadsheet of {len(usable)} file(s), {len(rows)} row(s), columns: {', '.join(cols)}, in "
            f"{time.monotonic() - started:.0f} s. Saved: {out}" + (" and opened it" if open_after else "") + "."
            + (f" {moved}" if moved else "")
            + (f" Skipped (nothing readable): {', '.join(skipped[:8])}." if skipped else "")
            + (f" {failed} file(s) could not be read." if failed else "")
            + (f"\nFirst rows: {preview}" if preview else ""))


# ---------------------------------------------------------------- edit_spreadsheet

EDITABLE = {".xlsx", ".xlsm"}
_TOTAL = re.compile(r"\s*(grand\s+)?totals?\b", re.I)
_SET = re.compile(r"\s*([A-Za-z]{1,3}[1-9][0-9]*)\s*(?:=|:|->)\s*(.*)$", re.S)
_SHOW = 40                      # rows shown back


def _cells(raw) -> list[str]:
    """One new row: a list, or text with the cells separated by '|' (or tabs)."""
    if isinstance(raw, (list, tuple)):
        return ["" if c is None else str(c) for c in raw]
    text = str(raw or "")
    return [c.strip() for c in (text.split("\t") if "\t" in text else text.split("|"))]


def _typed(text: str):
    """A typed-in cell: '=SUM(B2:B5)' stays a formula; numbers, amounts, percentages and dates get their type."""
    from mint.tools.convert import _value
    text = text.strip()
    if text.startswith("="):
        return text, None
    return _value(text)


def _number(value) -> float | None:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _shown(sheet, first: int, last: int, sums: dict[str, float]) -> str:
    lines = []
    for r in range(first, last + 1):
        cells = []
        for c in sheet[r]:
            value = c.value
            if isinstance(value, str) and value.startswith("=") and c.coordinate in sums:
                value = f"{value} (= {sums[c.coordinate]:g})"
            elif isinstance(value, dt.datetime):
                value = value.strftime("%Y-%m-%d %H:%M")
            elif isinstance(value, dt.date):
                value = value.isoformat()
            cells.append("" if value is None else str(value))
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            lines.append(f"{r}: " + " | ".join(cells))
    return "\n".join(lines)


def _extent(ws) -> tuple[int, int, int]:
    """(first used row, last used row, first used column) of a sheet; (1, 0, 1) when it is empty."""
    used = [(c.row, c.column) for r in ws.iter_rows() for c in r if c.value is not None and str(c.value).strip()]
    if not used:
        return 1, 0, 1
    rows = [r for r, _c in used]
    return min(rows), max(rows), min(c for _r, c in used)


def _create_if_new(path: str, changing: bool) -> None:
    """A new spreadsheet the user asked for ("make a budget in GenOffice and save it as budget"): an empty .xlsx where
    they said, so its rows go in through this tool - exact - instead of typing into a spreadsheet app's grid."""
    from pathlib import Path
    if not changing or not path.lower().endswith(".xlsx"):
        return
    file = Path(path).expanduser()
    if file.exists() or not file.parent.is_dir():
        return
    from openpyxl import Workbook
    Workbook().save(file)


def edit_spreadsheet(path: str, add_rows=None, set_cells=None, total: bool = False, sheet: str = "") -> str:
    """Change an existing .xlsx in place (a backup is kept): add rows at the bottom, set cells, add or refresh a
    Total row of =SUM formulas. With no change asked for, show what is in it."""
    import shutil
    import zipfile
    from copy import copy

    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from mint.tools import harness as harness_tools
    from mint.tools import undo
    _create_if_new(str(path or ""), bool(add_rows or set_cells or total))
    file, why = harness_tools._resolve(str(path or ""))
    if file is None:
        return f"FAILED: {why}"
    if file.suffix.lower() not in EDITABLE:
        return (f"FAILED: edit_spreadsheet changes .xlsx files; {file.name} is not one"
                + (" (a CSV is plain text: use write_file)." if file.suffix.lower() in {".csv", ".tsv"} else "."))
    why = harness_tools._blocked(file, write=True)
    if why:
        return f"FAILED: {why}"
    rows_in = [_cells(r) for r in (add_rows or []) if any(c.strip() for c in _cells(r))]
    sets = [str(s) for s in (set_cells or []) if str(s).strip()]
    changing = bool(rows_in or sets or total)
    try:
        with zipfile.ZipFile(file) as archive:
            parts = archive.namelist()
    except zipfile.BadZipFile:
        return f"FAILED: {file.name} is not a readable .xlsx (it may be damaged or password protected)."
    extras = [what for key, what in (("xl/charts/", "charts"), ("xl/pivotTables/", "pivot tables"),
                                     ("xl/drawings/", "pictures or shapes")) if any(p.startswith(key) for p in parts)]
    if changing and extras:
        return (f"FAILED: {file.name} has {' and '.join(extras)}, which editing it here would lose. Nothing was "
                "changed; the user can make this change in Excel or Numbers.")
    book = load_workbook(file, keep_vba=file.suffix.lower() == ".xlsm")
    if sheet:
        match = next((ws for ws in book.worksheets if ws.title.lower() == sheet.strip().lower()), None)
        if match is None:
            return f"FAILED: {file.name} has no sheet '{sheet}'; its sheets: {', '.join(book.sheetnames)}."
        ws = match
    else:
        ws = book.active
    if not changing:
        first, last, _left = _extent(ws)
        more = f"\n… {last - first + 1 - _SHOW} more rows" if last - first + 1 > _SHOW else ""
        return (f"{harness_tools._short(file)}, sheet '{ws.title}' (sheets: {', '.join(book.sheetnames)}), rows "
                f"{first}-{last}:\n{_shown(ws, first, min(last, first + _SHOW - 1), {})}{more}")

    done = []
    for text in sets:
        found = _SET.match(text)
        if not found:
            return f"FAILED: '{text}' is not a cell change; write it like 'B3=250'. Nothing was changed."
        value, shape = _typed(found.group(2))
        cell = ws[found.group(1).upper()]
        cell.value = value
        if shape:
            cell.number_format = shape
        done.append(f"set {cell.coordinate}")
    first, last, left = _extent(ws)

    # An existing Total row at the bottom comes off and goes back below the new rows, summing them too.
    total_row = last if last > first and _TOTAL.match(str(ws.cell(last, left).value or "")) else 0
    label = str(ws.cell(total_row, left).value).strip() if total_row else "Total"
    if total_row and rows_in:
        ws.delete_rows(total_row)
        last -= 1
    template = max(last, 1)
    for cells in rows_in:
        last += 1
        for i, text in enumerate(cells):
            cell = ws.cell(last, left + i)
            above = ws.cell(template, left + i)
            if above.has_style:
                cell._style = copy(above._style)
            value, shape = _typed(text) if text.strip() else (None, None)
            cell.value = value
            if shape:
                cell.number_format = shape
    if rows_in:
        done.append(f"added {len(rows_in)} row{'s' if len(rows_in) > 1 else ''}")

    sums: dict[str, float] = {}
    if total or (total_row and rows_in):
        if total_row and not rows_in:
            last -= 1                                 # refresh the Total row that is there
        header = not any(_number(c.value) is not None for c in ws[first][left:])
        top = first + 1 if header else first
        target = last + 1
        wrote = []
        width = max((c.column for r in ws.iter_rows(min_row=first, max_row=last) for c in r
                     if c.value is not None and str(c.value).strip()), default=left)
        for column in range(left + 1, width + 1):
            filled = [v for v in (ws.cell(r, column).value for r in range(top, last + 1))
                      if v is not None and str(v).strip()]
            numbers = [n for n in map(_number, filled) if n is not None]
            formulas = [v for v in filled if isinstance(v, str) and v.startswith("=")]
            if not numbers or len(numbers) + len(formulas) < len(filled) / 2:
                continue
            letter = get_column_letter(column)
            cell = ws.cell(target, column, f"=SUM({letter}{top}:{letter}{last})")
            above = ws.cell(last, column)
            if above.has_style:
                cell._style = copy(above._style)
            cell.font = Font(bold=True)
            if not formulas:                          # its value, to say back (formulas are worked out on opening)
                sums[cell.coordinate] = sum(numbers)
            wrote.append(f"{letter}: {cell.value}" + (f" = {sum(numbers):g}" if not formulas else ""))
        if not wrote:
            return f"FAILED: no column of {file.name} (sheet '{ws.title}') holds numbers to total. Nothing was changed."
        head = ws.cell(target, left, label)
        head.font = Font(bold=True)
        last = target
        done.append("a Total row (" + "; ".join(wrote) + ")")

    harness_tools.BACKUPS.mkdir(parents=True, exist_ok=True)
    backup, n = harness_tools.BACKUPS / f"{time.strftime('%Y%m%d-%H%M%S')}-{file.name}", 2
    while backup.exists():                            # two edits in one second keep both versions
        backup, n = backup.with_name(f"{time.strftime('%Y%m%d-%H%M%S')}-{n}-{file.name}"), n + 1
    shutil.copy2(file, backup)
    try:
        book.save(file)
    except OSError as error:
        return f"FAILED: could not save {harness_tools._short(file)}: {error.strerror or error}"
    undo.record("file", f"editing {harness_tools._short(file)}",
                {"kind": "file_restore", "path": str(file), "backup": str(backup), "mtime": file.stat().st_mtime})
    start = max(first, last - _SHOW + 1)
    return (f"Edited {harness_tools._short(file)} (sheet '{ws.title}'): {', '.join(done)}. The previous version is "
            f"saved at {harness_tools._short(backup)}. If it is open in Numbers or Excel, close that window without "
            f"saving and open it again to see the change. It now reads:\n{_shown(ws, start, last, sums)}")


PROMPT = """Changing an existing .xlsx: call edit_spreadsheet straight away, even if the file is open in \
Numbers or Excel - never try to close or quit apps first. Spreadsheets from files: "put my invoices / receipts / CVs / statements into a spreadsheet", "make a \
table of these PDFs" -> make_spreadsheet with the folder (or files) and what the rows are; pass columns only if \
the user names them. It reads every file (scans too) and saves an .xlsx in Mint's Spreadsheets folder. Changing an \
existing .xlsx ("add a row to budget.xlsx", "put a total at the bottom", "set B3 to 250") -> edit_spreadsheet with \
its path: it changes the file itself (a backup is kept). Do not open it in Numbers or Excel and type into cells - \
Numbers does not save back to .xlsx. edit_spreadsheet with only the path shows what is in it."""


def declarations():
    from google.genai import types
    S = types.Type.STRING
    return [types.FunctionDeclaration(
        name="make_spreadsheet",
        description=("Read many documents (PDFs, Word, text, scans, photos of receipts) and put the same fields from "
                     "each into one .xlsx spreadsheet, one row per document (or per item in a statement), then open "
                     "it. For invoices, receipts, CVs, bank statements, contracts, forms."),
        parameters=types.Schema(type=types.Type.OBJECT, properties={
            "source": types.Schema(type=S, description="a folder ('~/Downloads', 'desktop/receipts'), a file, or "
                                                       "several paths separated by ';'"),
            "pattern": types.Schema(type=S, description="optional file pattern, e.g. '*.pdf' or '*invoice*'"),
            "what": types.Schema(type=S, description="what the files are / what each row is, e.g. 'invoices'"),
            "columns": types.Schema(type=types.Type.ARRAY, items=types.Schema(type=S),
                                    description="only if the user named them"),
            "name": types.Schema(type=S, description="file name for the spreadsheet (optional)"),
            "save_to": types.Schema(type=S, description="where the user asked for it: a folder or a full .xlsx "
                                                        "path; empty = Mint's Spreadsheets folder")},
            required=["source"])),
        types.FunctionDeclaration(
            name="edit_spreadsheet",
            description=("Make or change an Excel .xlsx file without typing into an app (a new path makes a new file): add rows at the bottom, set "
                         "cells, add a Total row (=SUM of every number column, below the data; an existing Total row "
                         "moves below new rows). A backup is kept. With only the path, shows what is in it. Use this "
                         "instead of typing into Numbers or Excel."),
            parameters=types.Schema(type=types.Type.OBJECT, properties={
                "path": types.Schema(type=S, description="the .xlsx file, e.g. '~/Documents/budget.xlsx'"),
                "add_rows": types.Schema(type=types.Type.ARRAY, items=types.Schema(type=S),
                                         description="rows to add at the bottom, one string per row with the cells "
                                                     "in column order separated by ' | ', e.g. 'Travel | 1200'"),
                "set_cells": types.Schema(type=types.Type.ARRAY, items=types.Schema(type=S),
                                          description="cells to change, e.g. 'B3=250', 'A7=Notes', 'C9==B9*2' "
                                                      "(a formula)"),
                "total": types.Schema(type=types.Type.BOOLEAN,
                                      description="add a Total row at the bottom that sums the number columns"),
                "sheet": types.Schema(type=S, description="which sheet (default: the one that opens first)")},
                required=["path"]))]


def tool(args: dict) -> str:
    """Up to a handful of files: answered now. More: made in the background, with a message when done."""
    import threading
    from mint.tools import saveto
    # Worked out now, while the request is at hand: a big job runs in the background, later.
    save_to = str(args.get("save_to") or "") or saveto.requested(".xlsx")
    job = (str(args.get("source") or ""), str(args.get("what") or ""), args.get("columns"),
           str(args.get("pattern") or ""), str(args.get("name") or ""), True, save_to)
    count = len(_files(job[0], job[3]))
    if count <= 6:
        return make_spreadsheet(*job)

    def run():
        try:
            result = make_spreadsheet(*job)
        except Exception as error:
            log.exception("spreadsheet failed")
            result = f"FAILED: {error}"
        try:
            from mint.tools.work import _notify_mint
            _notify_mint(f"(Mint's spreadsheet maker, not the user.) {result} Tell the user briefly.")
        except Exception:
            pass
    threading.Thread(target=run, daemon=True, name="spreadsheet").start()
    return (f"Reading {count} files into a spreadsheet in the background (about {max(1, count // 6 * 15 // 60)} "
            "min). A message comes when it is ready; tell the user and carry on.")


def edit_tool(args: dict) -> str:
    return edit_spreadsheet(str(args.get("path") or ""), args.get("add_rows"), args.get("set_cells"),
                            args.get("total") is True, str(args.get("sheet") or ""))


HANDLERS = {"make_spreadsheet": tool, "edit_spreadsheet": edit_tool}
