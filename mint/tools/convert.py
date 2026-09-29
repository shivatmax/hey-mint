"""Text, tables and documents into the shape the user needs.

"copy the text on my screen", "OCR this image and copy it", "make an Excel of this table",
"convert this English PDF to a Hindi Word document", "turn report.docx into a PDF".

ocr_copy reads the text of the front window, the whole screen, an image or PDF, or a
picture on the clipboard with the Mac's own text recognition (Vision, any script, the
language found by itself - nothing is uploaded), puts it back together in reading order
(side-by-side pieces on one row, paragraphs joined, or the exact lines when asked) and
copies it to the clipboard.

data_to_sheet turns a table into a real .xlsx: from the screen (Gemini reads the
picture), a PDF, image, Word or text file, a CSV, or what is on the clipboard (copied
cells are read here, without Gemini). Numbers become numbers and dates dates, with a bold
frozen header, in the Spreadsheets folder of Mint's storage.

convert_document rebuilds a document (PDF, Word, RTF, text, Markdown, HTML, or a scan)
as headings, paragraphs, lists and tables - PDFs from their text and the size and spacing
of each line, scanned pages by text recognition - optionally translates it with Gemini a
few pages at a time (the structure is kept; Indian and other scripts get a font that has
them), and writes Word, PDF, Markdown, text or HTML next to the original, e.g.
"report (Hindi).docx". Anything taking over 20 seconds finishes in the background with a
message when it is done; convert_document status=true tells how far it has got.
"""

from __future__ import annotations

import csv
import datetime as dt
import html
import io
import json
import logging
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from html.parser import HTMLParser
from pathlib import Path

log = logging.getLogger("mint.tools.convert")

# Flash Lite first: in testing it read tables and translated a 6,000-character piece in 3-6 s, while the Flash
# models were often busy (503) or out of quota (429) and took 20-30 s when they answered.
MODELS = ["gemini-3.5-flash-lite", "gemini-3.5-flash", "gemini-3.7-flash", "gemini-3.1-flash-lite"]
IMAGES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".heic": "image/heic",
          ".webp": "image/webp", ".gif": "image/gif", ".tiff": "image/tiff", ".tif": "image/tiff",
          ".bmp": "image/bmp"}
RICH = {".docx", ".doc", ".rtf", ".rtfd", ".odt", ".wordml", ".webarchive"}      # textutil reads these
PLAIN = {".txt", ".text", ".md", ".markdown", ".html", ".htm", ".csv", ".tsv"}
SHEET_SOURCES = {".pdf", ".json"} | PLAIN | RICH | set(IMAGES)   # what data_to_sheet can read a table from
FORMATS = ("docx", "pdf", "md", "txt", "html", "rtf")
CHUNK = 6000                  # characters of HTML per translation request
MAX_INLINE = 18_000_000       # Gemini's inline limit is 20 MB per request
WAIT = 20                     # seconds to answer in before finishing in the background

# language -> (code, a font on every Mac that has its script, right to left)
SCRIPTS = {
    "hindi": ("hi", "Kohinoor Devanagari", False), "marathi": ("mr", "Kohinoor Devanagari", False),
    "nepali": ("ne", "Kohinoor Devanagari", False), "sanskrit": ("sa", "Kohinoor Devanagari", False),
    "konkani": ("kok", "Kohinoor Devanagari", False), "maithili": ("mai", "Kohinoor Devanagari", False),
    "bengali": ("bn", "Kohinoor Bangla", False), "bangla": ("bn", "Kohinoor Bangla", False),
    "assamese": ("as", "Kohinoor Bangla", False), "gujarati": ("gu", "Kohinoor Gujarati", False),
    "punjabi": ("pa", "Mukta Mahee", False), "tamil": ("ta", "Tamil Sangam MN", False),
    "telugu": ("te", "Kohinoor Telugu", False), "kannada": ("kn", "Noto Sans Kannada", False),
    "malayalam": ("ml", "Malayalam Sangam MN", False), "odia": ("or", "Oriya Sangam MN", False),
    "oriya": ("or", "Oriya Sangam MN", False), "sinhala": ("si", "Sinhala Sangam MN", False),
    "urdu": ("ur", "Noto Nastaliq Urdu", True), "arabic": ("ar", "Geeza Pro", True),
    "persian": ("fa", "Geeza Pro", True), "farsi": ("fa", "Geeza Pro", True), "hebrew": ("he", "Arial Hebrew", True),
    "japanese": ("ja", "Hiragino Sans", False), "chinese": ("zh", "PingFang SC", False),
    "korean": ("ko", "Apple SD Gothic Neo", False), "thai": ("th", "Thonburi", False),
    "burmese": ("my", "Myanmar Sangam MN", False), "khmer": ("km", "Khmer Sangam MN", False),
    "lao": ("lo", "Lao Sangam MN", False)}
LATIN_FONT = "Helvetica Neue"

_jobs: dict[str, dict] = {}   # what each running (or just finished) job is doing, for status and the island
_jobs_lock = threading.Lock()


# ---------------------------------------------------------------- where things come from

def _file(raw: str, same_name: set[str] | None = None) -> tuple[Path | None, str]:
    """The file the user means: a path, a name Mint can find, or "this" (selected in Finder or open in the
    front app). With `same_name` (suffixes): a missing file whose folder holds exactly one file of the same
    name with one of those suffixes ("scores.pdf" when only scores.csv is there) gives that file, and the
    second value says so."""
    text = str(raw or "").strip().strip("\"'")
    if text.lower() in {"", "this", "it", "selected", "selection", "this file", "the selected file"}:
        from mint.tools.clipboard import _this_path
        found, _what = _this_path()
        first = found.splitlines()[0].strip() if found else ""
        if not first or not Path(first).is_file():
            return None, "Which file? Say its name, or select it in Finder first."
        return Path(first), ""
    from mint.tools.harness import _resolve
    path, why = _resolve(text)
    if path is None:
        if same_name:
            return _same_name(text, why, same_name)
        return None, why
    if not path.is_file():
        return None, f"{path} is not a file."
    return path, ""


def _same_name(text: str, why: str, suffixes: set[str]) -> tuple[Path | None, str]:
    """The user named a file that is not there: the one file next to where it would be with the same name
    and another (readable) extension, else (None, why) - listing them when there are several."""
    from mint.tools.harness import _blocked, _resolve, _short
    if "/" not in text:                   # a bare name: no folder the user meant to look next to
        return None, why
    try:
        meant, _ = _resolve(text, must_exist=False)
    except Exception:
        meant = None
    if meant is None or meant.exists() or not meant.parent.is_dir() or _blocked(meant):
        return None, why
    others = sorted(p for p in meant.parent.iterdir() if p.is_file() and p != meant
                    and p.stem.lower() == meant.stem.lower() and p.suffix.lower() in suffixes)
    if len(others) == 1:
        return others[0], (f"There is no {meant.name} in {_short(meant.parent)}; used {others[0].name}, the file "
                           "with that name there - tell the user.")
    if others:
        return None, (f"{why} That folder has {', '.join(p.name for p in others)} - ask the user which one, "
                      "then pass its path.")
    return None, why


def _grab(source: str):
    """A picture of the front window or the whole main display (PIL image)."""
    if source == "screen":
        from mint.screen.ocr import _screen
        image = _screen()[0]
    else:
        from mint.tools.translate import _capture
        image = _capture()[0]
    low, high = image.convert("L").resize((64, 40)).getextrema()
    if high - low < 12:
        raise RuntimeError("the capture came back blank - Mint needs Screen Recording permission")
    return image


def _clipboard(prefer_image: bool = False) -> tuple[str, object]:
    """What is on the clipboard: ("file", Path), ("image", PIL image), ("text", str) or ("", None). Copied cells
    come as text and a picture of them: the text wins unless `prefer_image`."""
    from mint.tools import clipboard as clip_tools
    board = clip_tools._board()
    kinds = set(clip_tools._types(board))
    if kinds & clip_tools._CONCEALED:
        return "", None
    files = clip_tools._files(board)
    if files and Path(files[0]).is_file():
        return "file", Path(files[0])
    image = clip_tools._image(board)
    if image is not None and (prefer_image or not board.stringForType_("public.utf8-plain-text")):
        import PIL.Image
        return "image", PIL.Image.open(io.BytesIO(bytes(image.TIFFRepresentation()))).convert("RGB")
    text = board.stringForType_("public.utf8-plain-text")
    return ("text", str(text)) if text else ("", None)


# ---------------------------------------------------------------- reading text in reading order

def _recognise(image) -> list[dict]:
    """Every line of text in a picture ({"text", "box"} in its pixels), by the Mac's text recognition."""
    from mint.tools.translate import _lines
    return _lines(image.convert("RGB"))


def _pdf_lines(page) -> list[dict]:
    """The text of one PDF page as pieces of lines with their boxes (top-down, in points). A line is cut where
    its words are far apart, so the cells of a table stay apart."""
    import Quartz
    box = page.boundsForBox_(Quartz.kPDFDisplayBoxMediaBox)
    text = str(page.string() or "")
    units = [0]                                     # where each character starts in UTF-16, as PDFKit counts
    for ch in text:
        units.append(units[-1] + (2 if ord(ch) > 0xFFFF else 1))
    pieces: list[dict] = []
    for match in re.finditer(r"\S+", text):
        start, end = units[match.start()], units[match.end()]
        b = page.selectionForRange_((start, end - start)).boundsForPage_(page)
        top = box.origin.y + box.size.height - (b.origin.y + b.size.height)
        word = (b.origin.x, top, b.origin.x + b.size.width, top + b.size.height)
        last = pieces[-1] if pieces else None
        if last is not None:
            x0, y0, x1, y1 = last["box"]
            height = max(1.0, y1 - y0, word[3] - word[1])
            beside = abs((word[1] + word[3]) / 2 - (y0 + y1) / 2) < 0.5 * height
            if b.size.height <= 0 or (beside and -0.5 * height < word[0] - x1 < 0.7 * height):
                last["end"] = match.end()
                if b.size.height > 0:
                    last["box"] = (min(x0, word[0]), min(y0, word[1]), max(x1, word[2]), max(y1, word[3]))
                continue
        pieces.append({"start": match.start(), "end": match.end(), "box": word})
    return [{"text": " ".join(text[piece["start"]:piece["end"]].split()), "box": piece["box"]} for piece in pieces]


def _page_image(page, longest: int = 2400):
    """One PDF page drawn as a picture, for reading scans."""
    import PIL.Image
    import Quartz
    box = page.boundsForBox_(Quartz.kPDFDisplayBoxMediaBox)
    scale = min(3.0, longest / max(box.size.width, box.size.height, 1))
    width, height = int(box.size.width * scale), int(box.size.height * scale)
    space = Quartz.CGColorSpaceCreateDeviceRGB()
    context = Quartz.CGBitmapContextCreate(None, width, height, 8, width * 4, space,
                                           Quartz.kCGImageAlphaPremultipliedLast)
    Quartz.CGContextSetRGBFillColor(context, 1, 1, 1, 1)
    Quartz.CGContextFillRect(context, ((0, 0), (width, height)))
    Quartz.CGContextScaleCTM(context, scale, scale)
    Quartz.CGContextTranslateCTM(context, -box.origin.x, -box.origin.y)
    page.drawWithBox_toContext_(Quartz.kPDFDisplayBoxMediaBox, context)
    cg = Quartz.CGBitmapContextCreateImage(context)
    data = Quartz.CGDataProviderCopyData(Quartz.CGImageGetDataProvider(cg))
    row = Quartz.CGImageGetBytesPerRow(cg)
    return PIL.Image.frombuffer("RGBA", (width, height), bytes(data), "raw", "RGBA", row, 1).convert("RGB")


def _rows(lines: list[dict]) -> list[dict]:
    """Lines side by side (the cells of a table, pieces of one line) into rows, top to bottom."""
    rows: list[dict] = []
    for line in sorted(lines, key=lambda item: (item["box"][1] + item["box"][3]) / 2):
        x0, y0, x1, y1 = line["box"]
        middle, height = (y0 + y1) / 2, max(1.0, y1 - y0)
        if rows and abs(middle - rows[-1]["middle"]) < 0.45 * min(height, rows[-1]["height"]):
            row = rows[-1]
            row["cells"].append(line)
            row["top"], row["bottom"] = min(row["top"], y0), max(row["bottom"], y1)
        else:
            rows.append({"middle": middle, "height": height, "top": y0, "bottom": y1, "cells": [line]})
    for row in rows:
        row["cells"].sort(key=lambda item: item["box"][0])
    return rows


_BULLET = re.compile(r"^\s*([•●▪◦‣∙·○■□➢►–-]|\*|\(?\d{1,2}[.)]|\(?[a-zA-Z][.)])\s+")


def _blocks(lines: list[dict]) -> list[dict]:
    """Lines with boxes -> headings, paragraphs, list items and tables in reading order:
    [{"kind": "h1" | "h2" | "p" | "li" | "table", "lines": [text, ...] or "rows": [[cell, ...], ...]}].
    Headings are the taller lines, paragraphs end at a wider gap, list items start with a bullet or number or
    are indented (bullets drawn as pictures are not text), tables are rows of cells far apart."""
    rows = _rows(lines)
    if not rows:
        return []
    heights = sorted(row["height"] for row in rows)
    body = heights[len(heights) // 2]
    steps = sorted(b["top"] - a["top"] for a, b in zip(rows, rows[1:]) if 0 < b["top"] - a["top"] < 4 * body)
    step = steps[len(steps) // 4] if steps else body * 1.3        # the usual distance from one line to the next
    starts = [round(row["cells"][0]["box"][0] / body) for row in rows]
    left = max(set(starts), key=starts.count) * body              # the page's usual left edge
    ends = sorted(row["cells"][-1]["box"][2] for row in rows)
    right = ends[int(len(ends) * 0.9)]                            # and right edge
    blocks: list[dict] = []
    previous = None
    for row in rows:
        cells = [cell["text"].strip() for cell in row["cells"]]
        apart = any(b["box"][0] - a["box"][2] > 0.6 * body for a, b in zip(row["cells"], row["cells"][1:]))
        text = "  ".join(cells)
        x = row["cells"][0]["box"][0]
        indented = left + 1.2 * body < x < left + 6 * body
        if len(cells) > 1 and apart:
            kind = "cells"
        elif row["height"] >= 1.55 * body and len(text) < 160:
            kind = "h1"
        elif row["height"] >= 1.2 * body and len(text) < 160:
            kind = "h2"
        elif _BULLET.match(text):
            kind = "li"
        else:
            kind = "p"
        distance = row["top"] - previous["top"] if previous else 0
        gap = previous is None or distance > 1.45 * max(step, previous["height"])
        wrapped = previous is not None and previous["cells"][-1]["box"][2] > right - 4 * body
        last = blocks[-1] if blocks else None
        if kind == "cells":
            if last and last["kind"] == "table" and distance < 3.5 * max(step, body):
                last["rows"].append(cells)
            else:
                blocks.append({"kind": "table", "rows": [cells]})
        elif last and not gap and last["kind"] in ("li", "indent") and kind == "p" and wrapped \
                and abs(x - last["x"]) < 0.8 * body:
            last["lines"].append(text)                               # a list item, wrapped
        elif last and not gap and last["kind"] == "indent" and kind == "p" and wrapped and x < last["x"]:
            last["kind"] = "p"                                       # a paragraph with its first line indented
            last["lines"].append(text)
        elif last and not gap and kind == last["kind"] and kind in ("p", "h1", "h2") \
                and not (kind == "p" and indented):
            last["lines"].append(text)                               # the same paragraph, wrapped
        else:
            blocks.append({"kind": "indent" if kind == "p" and indented else kind, "lines": [text], "x": x})
        previous = row
    for block in blocks:
        if block["kind"] == "table" and len(block["rows"]) < 2:      # a lone row of cells is just a line
            block.update(kind="p", lines=["  ".join(block.pop("rows")[0])])
    for i, block in enumerate(blocks):                               # indented lines next to others: a list
        if block["kind"] == "indent":
            near = [b["kind"] for b in blocks[max(0, i - 1):i + 2] if b is not block]
            block["kind"] = "li" if any(k in ("li", "indent") for k in near) else "p"
    return blocks


def _no_space(a: str, b: str) -> bool:
    """Scripts written without spaces between words (Chinese, Japanese, Thai)."""
    def dense(ch):
        return "\u3000" <= ch <= "\u9fff" or "\uff00" <= ch <= "\uffef" or "\u0e00" <= ch <= "\u0e7f"
    return bool(a and b) and dense(a[-1]) and dense(b[0])


def _join(lines: list[str]) -> str:
    """Wrapped lines back into one paragraph (a word broken with a hyphen is put back together)."""
    out = ""
    for line in lines:
        if not out:
            out = line
        elif re.search(r"[a-z]-$", out) and line[:1].islower():
            out = out[:-1] + line
        else:
            out += ("" if _no_space(out, line) else " ") + line
    return out


def _as_text(blocks: list[dict], exact_lines: bool = False) -> str:
    """Blocks as plain text: a blank line between paragraphs, list items one under another, table cells
    separated by tabs (so they paste into a spreadsheet as cells)."""
    out = ""
    for i, block in enumerate(blocks):
        if block["kind"] == "table":
            text = "\n".join("\t".join(row) for row in block["rows"])
        elif block["kind"] == "raw":
            text = block["lines"][0]
        else:
            text = "\n".join(block["lines"]) if exact_lines else _join(block["lines"])
            if block["kind"] == "li" and not _BULLET.match(text):
                text = "• " + text
        together = block["kind"] == "li" and i and blocks[i - 1]["kind"] == "li"
        out += ("\n" if together else "\n\n" if out else "") + text
    return out


def _as_html(blocks: list[dict]) -> str:
    out, in_list = [], False
    for block in blocks:
        if in_list and block["kind"] != "li":
            out.append("</ul>")
            in_list = False
        if block["kind"] == "table":
            rows = block["rows"]
            width = max(len(row) for row in rows)
            cells = ["<tr>" + "".join(f"<{tag}>{html.escape(cell)}</{tag}>" for cell in row + [""] * (width - len(row)))
                     + "</tr>" for i, row in enumerate(rows) for tag in ["th" if i == 0 else "td"]]
            out.append("<table>" + "".join(cells) + "</table>")
        elif block["kind"] == "li":
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{html.escape(_BULLET.sub('', _join(block['lines']), count=1))}</li>")
        else:
            out.append(f"<{block['kind']}>{html.escape(_join(block['lines']))}</{block['kind']}>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)


def _open_pdf(path: Path):
    import Quartz
    from Foundation import NSURL
    document = Quartz.PDFDocument.alloc().initWithURL_(NSURL.fileURLWithPath_(str(path)))
    if document is None:
        raise RuntimeError(f"{path.name} could not be opened as a PDF")
    if document.isLocked():
        raise RuntimeError(f"{path.name} is protected with a password")
    return document


def _pdf_blocks(path: Path, job: dict | None = None) -> tuple[list[list[dict]], int]:
    """Each page of a PDF as blocks: its text layer, or (a scan) text recognition of the page picture.
    -> (pages, how many pages were read from pictures)."""
    document = _open_pdf(path)
    pages, scanned = [], 0
    for index in range(document.pageCount()):
        page = document.pageAtIndex_(index)
        lines = _pdf_lines(page)
        if sum(len(line["text"]) for line in lines) < 20:
            lines = _recognise(_page_image(page))
            scanned += 1 if lines else 0
        pages.append(_blocks(lines))
        if job is not None:
            job.update(step=f"reading page {index + 1} of {document.pageCount()}", done=index + 1,
                       total=document.pageCount())
    return pages, scanned


def _image_file(path: Path):
    import PIL.Image
    if path.suffix.lower() == ".heic":            # PIL cannot read HEIC; sips can
        work = Path(tempfile.mkdtemp(prefix="mint-convert-"))
        subprocess.run(["sips", "-s", "format", "png", str(path), "--out", str(work / "image.png")],
                       capture_output=True, timeout=60, check=False)
        image = PIL.Image.open(work / "image.png").convert("RGB")
        shutil.rmtree(work, ignore_errors=True)
        return image
    return PIL.Image.open(path).convert("RGB")


def _textutil(path: Path, kind: str) -> str:
    done = subprocess.run(["textutil", "-convert", kind, "-stdout", str(path)], capture_output=True, timeout=120,
                          check=False)
    if done.returncode != 0:
        raise RuntimeError(f"could not read {path.name}: {done.stderr.decode(errors='replace').strip()[:200]}")
    return done.stdout.decode("utf-8", errors="replace")


# ---------------------------------------------------------------- ocr_copy

def _copy(text: str) -> None:
    from mint.tools.clipboard import _put_text
    _put_text(text)


SHOWN = 3000          # up to this many characters of recognised text go back to the model in full


def ocr_copy(source: str = "window", path: str = "", exact_lines: bool = False, job: dict | None = None,
             save_to: str = "") -> str:
    started = time.monotonic()
    what, pages = "", []
    try:
        if path or source == "file":
            file, why = _file(path)
            if file is None:
                return why
            what, pages = _file_blocks(file, job)
        elif source == "clipboard_image":
            kind, value = _clipboard(prefer_image=True)
            if kind == "file":
                what, pages = _file_blocks(value, job)
            elif kind == "image":
                what, pages = "the picture on the clipboard", [_blocks(_recognise(value))]
            else:
                return "There is no picture (or image or PDF file) on the clipboard to read."
        else:
            what = "the whole screen" if source == "screen" else "the front window"
            pages = [_blocks(_recognise(_grab(source)))]
    except Exception as error:
        log.info("ocr_copy: %s", error)
        return f"FAILED: could not read {what or 'it'}: {error}"
    text = "\n\n".join(t for t in (_as_text(blocks, exact_lines) for blocks in pages) if t.strip())
    if not text.strip():
        return f"Found no text in {what}."
    _copy(text)
    from mint.knowledge.skills import has_secret
    took = f"Copied {len(text):,} characters of text from {what} to the clipboard in {time.monotonic() - started:.0f} s."
    saved = ""
    if save_to.strip():
        from mint.tools.harness import write_file
        saved = " " + write_file({"path": save_to.strip(), "content": text.rstrip("\n") + "\n", "mode": "create"})
    if has_secret(text):
        return f"{took}{saved} (It looks like it holds a password or key, so it is not read out here.)"
    # The whole text, not just its start: a window's first lines are often only tabs and toolbars, and the
    # model may need the rest (to save it, answer from it, or check it).
    if len(text) <= SHOWN:
        return f"{took}{saved} The text:\n{text}"
    return (f"{took}{saved} The first {SHOWN:,} characters:\n{text[:SHOWN]}\n[... {len(text) - SHOWN:,} more "
            "characters are on the clipboard.]")


def _file_blocks(path: Path, job: dict | None = None) -> tuple[str, list[list[dict]]]:
    """(what it is, pages of blocks) for any readable file, for copying its text."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages, scanned = _pdf_blocks(path, job)
        return (f"{path.name} ({len(pages)} pages" + (f", {scanned} read from pictures" if scanned else "") + ")",
                pages)
    if suffix in IMAGES:
        return path.name, [_blocks(_recognise(_image_file(path)))]
    if suffix in RICH or suffix in {".html", ".htm"}:
        text = _textutil(path, "txt")
    else:
        text = path.read_text(errors="replace")
    text = re.sub(r"[ \t]+\n", "\n", text).strip()               # a text document's text, as it is
    return path.name, [[{"kind": "raw", "lines": [text]}]] if text else [[]]


# ---------------------------------------------------------------- data_to_sheet

def _split_table(text: str) -> list[list[str]] | None:
    """Rows of copied cells (tab separated, as from Excel, Numbers or a web table) or of a CSV, when the text
    is clearly a table; else None."""
    lines = [line for line in text.strip("\n").splitlines() if line.strip()]
    if len(lines) < 2:
        return None
    for delimiter in ("\t", ",", ";", "|"):
        if delimiter == "|":                          # a Markdown table: no outer bars, no --- line
            lines = [line.strip().strip("|") for line in lines if not re.fullmatch(r"[\s|:-]+", line)]
        rows = list(csv.reader(lines, delimiter=delimiter))
        widths = [len(row) for row in rows]
        common = max(set(widths), key=widths.count)
        if common >= 2 and widths.count(common) >= 0.8 * len(rows):
            return [[cell.strip() for cell in row] for row in rows]
    return None


def _extract_tables(content: list, what: str) -> list[dict]:
    from mint.core import llm
    focus = f" (the user wants: {what})" if what else ""
    prompt = (f"Extract the tabular data{focus} from this. Return JSON "
              '{"sheets": [{"name": "<short name for the table>", "columns": ["..."], "rows": [["...", ...], ...]}]}'
              " - one sheet per separate table, every row in order, exactly as shown: nothing invented, merged or "
              "summarised, no total rows added. Numbers as plain numbers (no thousands separators; a currency sign "
              "goes into the column name, e.g. 'Price (USD)'); percentages as written, e.g. '12.5%'; dates as "
              "YYYY-MM-DD when the date is certain, otherwise as written. Empty cells as null. A list of records "
              "that is not drawn as a table (contacts, transactions, search results) becomes a table too. "
              'Menus, buttons and page furniture are not data. No data at all: {"sheets": []}.')
    text, _model = llm.generate([*content, prompt], MODELS, json_mode=True)
    data = llm.parse_json(text)
    sheets = data.get("sheets") if isinstance(data, dict) else data if isinstance(data, list) else []
    out = []
    for sheet in sheets or []:
        if not isinstance(sheet, dict):
            continue
        columns = [str(c if c is not None else "").strip() for c in sheet.get("columns") or []]
        rows = []
        for row in sheet.get("rows") or []:
            if isinstance(row, dict):
                row = [row.get(c) for c in columns]
            if isinstance(row, list) and any(v not in (None, "") for v in row):
                rows.append(row)
        if columns or rows:
            out.append({"name": str(sheet.get("name") or "Data"), "columns": columns, "rows": rows})
    return out


def _table_from_rows(rows: list[list[str]], name: str) -> dict:
    return {"name": name, "columns": rows[0], "rows": rows[1:]}


class _HTMLTables(HTMLParser):
    """The <table>s of an HTML page as rows of cell text (innermost tables only: a table holding other tables
    is page layout, not data)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: list[dict] = []
        self._stack: list[dict] = []          # tables being read, innermost last
        self._cell: list[str] | None = None
        self._span = 1
        self._skip = 0                        # inside <script>/<style>

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag == "table":
            if self._stack:
                self._stack[-1]["nested"] = True
            self._stack.append({"rows": [], "caption": "", "nested": False, "in_caption": False})
        elif not self._stack:
            return
        elif tag == "tr":
            self._stack[-1]["rows"].append([])
        elif tag in ("td", "th"):
            table = self._stack[-1]
            if not table["rows"]:
                table["rows"].append([])
            self._cell = []
            try:
                self._span = max(1, min(int(dict(attrs).get("colspan") or 1), 50))
            except ValueError:
                self._span = 1
        elif tag == "caption":
            self._stack[-1]["in_caption"] = True
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        elif not self._stack:
            return
        elif tag in ("td", "th"):
            self._end_cell()
        elif tag == "caption":
            self._stack[-1]["in_caption"] = False
        elif tag == "table":
            self._end_cell()
            table = self._stack.pop()
            rows = [row for row in table["rows"] if any(cell for cell in row)]
            if not table["nested"] and len(rows) >= 2 and max(len(row) for row in rows) >= 2:
                self.tables.append({"caption": table["caption"].strip(), "rows": rows})

    def _end_cell(self):
        if self._cell is not None and self._stack:
            text = " ".join("".join(self._cell).split())
            self._stack[-1]["rows"][-1].extend([text] + [""] * (self._span - 1))
        self._cell, self._span = None, 1

    def handle_data(self, data):
        if self._skip or not self._stack:
            return
        if self._cell is not None:
            self._cell.append(data)
        elif self._stack[-1]["in_caption"]:
            self._stack[-1]["caption"] += data


def _html_tables(page: str, name: str) -> list[dict]:
    """The data tables of an HTML page, read as they are (no Gemini); [] when it has none."""
    parser = _HTMLTables()
    try:
        parser.feed(page)
        parser.close()
    except Exception as error:                        # a broken page: let Gemini read its text instead
        log.info("html tables: %s", error)
        return []
    out = []
    for n, table in enumerate(parser.tables[:20], 1):
        label = table["caption"] or (name if len(parser.tables) == 1 else f"{name} {n}")
        out.append(_table_from_rows(table["rows"], label[:31] or "Data"))
    return out


_URL = re.compile(r"(https?://|www\.)\S+$", re.I)


def _tables_from_url(url: str, what: str) -> tuple[str, list[dict]]:
    """(what it was, tables) from a web page, fetched here - the page does not need to be open. Its <table>s
    are read as they are; a page without any (a list, or drawn by scripts) goes to Gemini as text."""
    import httpx

    from mint.agents.tools import UA
    if not url.lower().startswith(("http://", "https://")):
        url = "https://" + url
    response = httpx.get(url, timeout=25, headers={"User-Agent": UA}, follow_redirects=True)
    if response.status_code >= 400:
        raise RuntimeError(f"HTTP {response.status_code} for {url}")
    kind = response.headers.get("content-type", "").lower()
    page = response.text
    title = re.search(r"(?is)<title[^>]*>(.*?)</title>", page) if "html" in kind else None
    name = " ".join(html.unescape(title.group(1)).split())[:31] if title else "Data"
    label = f"the web page {url}"
    if "html" in kind:
        tables = _html_tables(page, name)
        if tables:
            return label, tables
        text = re.sub(r"(?is)<(script|style|noscript|svg)[^>]*>.*?</\1>", " ", _body(page)[1])
        return label, _extract_tables([text[:200_000]], what)
    if "text" in kind or "csv" in kind or "json" in kind:
        rows = _split_table(page)
        return label, [_table_from_rows(rows, name)] if rows else _extract_tables([page[:200_000]], what)
    raise RuntimeError(f"{url} is {kind or 'not a web page'}; download it and pass the file instead")


def _tables_from_file(path: Path, what: str) -> list[dict]:
    from google.genai import types
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        tables = _html_tables(path.read_text(errors="replace"), path.stem)
        if tables:
            return tables
    if suffix in {".csv", ".tsv"}:
        text = path.read_text(errors="replace")
        rows = list(csv.reader(text.splitlines(), delimiter="\t" if suffix == ".tsv" else
                               csv.Sniffer().sniff(text[:4000], ",;\t|").delimiter))
        return [_table_from_rows([r for r in rows if any(c.strip() for c in r)], path.stem)]
    size = path.stat().st_size
    if suffix in IMAGES:
        if suffix in {".png", ".jpg", ".jpeg", ".webp"} and size < MAX_INLINE:
            return _extract_tables([types.Part.from_bytes(data=path.read_bytes(), mime_type=IMAGES[suffix])], what)
        return _tables_from_image(_image_file(path), what)
    if suffix == ".pdf":
        if size < MAX_INLINE:
            return _extract_tables([types.Part.from_bytes(data=path.read_bytes(), mime_type="application/pdf")], what)
        pages, _scanned = _pdf_blocks(path)
        return _extract_tables(["\n\n".join(_as_text(p) for p in pages)[:200_000]], what)
    if suffix in RICH or suffix in {".html", ".htm"}:
        page = _textutil(path, "html") if suffix in RICH else path.read_text(errors="replace")
        return _extract_tables([_body(page)[1][:200_000]], what)
    text = path.read_text(errors="replace")
    rows = _split_table(text)
    return [_table_from_rows(rows, path.stem)] if rows else _extract_tables([text[:200_000]], what)


def _tables_from_image(image, what: str) -> list[dict]:
    from google.genai import types
    shot = image.copy()
    shot.thumbnail((2400, 2400))
    buffer = io.BytesIO()
    shot.convert("RGB").save(buffer, "JPEG", quality=90)
    return _extract_tables([types.Part.from_bytes(data=buffer.getvalue(), mime_type="image/jpeg")], what)


def _value(value):
    """A cell as the right type: numbers, percentages, amounts with a currency sign, dates."""
    from mint.tools.sheets import _cell
    if isinstance(value, bool) or value is None or isinstance(value, (int, float)):
        return value, None
    text = str(value).strip()
    if not text:
        return None, None
    plain = re.sub(r"(?<=[$€£₹¥])\s+|\s+(?=%$)|[,\u00a0\u202f]", "", text)      # "₹ 1,200" -> "₹1200"; "98200 11111" stays
    match = re.fullmatch(r"(-?)([$€£₹¥])?(-?\d+(?:\.\d+)?)(%)?", plain)
    if match and not (re.fullmatch(r"0\d+", match.group(3))):        # 007 and phone-like ids stay text
        number = float(match.group(3)) * (-1 if match.group(1) else 1)
        if match.group(4):
            return number / 100, "0.0#%"
        if match.group(2):
            return number, "#,##0.00"
        if len(match.group(3).replace(".", "").lstrip("-")) > 15:
            return text, None                                         # too long for a number (card, account)
        return (int(number) if "." not in match.group(3) else number), None
    for pattern, shape in ((r"\d{4}-\d{2}-\d{2}", "%Y-%m-%d"), (r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}", None)):
        if re.fullmatch(pattern, text):
            try:
                return (dt.datetime.strptime(text, shape).date() if shape
                        else dt.datetime.fromisoformat(text.replace(" ", "T"))), None
            except ValueError:
                pass
    return _cell(text), None


# Columns of numbers that are not amounts: kept as text so leading zeros and long digits survive.
_TEXT_COLUMN = re.compile(r"phone|mobile|\btel\b|whatsapp|\bfax\b|zip|pin ?code|postal|postcode|account|iban|"
                          r"ifsc|aadhaar|\bpan\b|card|\bid\b|passport", re.I)


def _add_total(path: Path) -> None:
    """A bold Total row under the data of the first sheet: =SUM of each number column."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    book = load_workbook(path)
    ws = book.worksheets[0]
    last = ws.max_row
    if last < 2:
        return
    ws.cell(last + 1, 1, "Total").font = Font(bold=True)
    for column in range(2, ws.max_column + 1):
        values = [ws.cell(r, column).value for r in range(2, last + 1)]
        if values and all(isinstance(v, (int, float)) for v in values if v is not None) and \
                any(isinstance(v, (int, float)) for v in values):
            letter = get_column_letter(column)
            cell = ws.cell(last + 1, column, f"=SUM({letter}2:{letter}{last})")
            cell.font = Font(bold=True)
            cell.number_format = ws.cell(last, column).number_format
    book.save(path)


def _write_book(path: Path, tables: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = Workbook()
    book.remove(book.active)
    used: set[str] = set()
    for table in tables:
        name = re.sub(r"[\[\]:*?/\\]", " ", table["name"]).strip()[:31] or "Data"
        base, n = name, 2
        while name.lower() in used:
            name, n = f"{base[:28]} {n}", n + 1
        used.add(name.lower())
        sheet = book.create_sheet(name)
        width = max([len(table["columns"])] + [len(row) for row in table["rows"]])
        columns = table["columns"] + [f"Column {i + 1}" for i in range(len(table["columns"]), width)]
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E6F4F1")
            cell.alignment = Alignment(vertical="center")
        as_text = {i for i, column in enumerate(columns, 1) if _TEXT_COLUMN.search(str(column))}
        for r, row in enumerate(table["rows"], 2):
            for c, raw in enumerate(list(row) + [None] * (width - len(row)), 1):
                value, shape = (None if raw in (None, "") else str(raw), None) if c in as_text else _value(raw)
                cell = sheet.cell(row=r, column=c, value=value)
                if shape:
                    cell.number_format = shape
                elif isinstance(value, dt.datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif isinstance(value, dt.date):
                    cell.number_format = "yyyy-mm-dd"
        for column in sheet.columns:
            longest = max(len(str(cell.value if cell.value is not None else "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(9, longest + 2))
        sheet.freeze_panes = "A2"
        if table["rows"]:
            sheet.auto_filter.ref = sheet.dimensions
    book.save(path)


def data_to_sheet(source: str = "window", path: str = "", what: str = "", name: str = "",
                  open_after: bool = True, job: dict | None = None, save_to: str = "", data: str = "") -> str:
    from mint.core import config
    from mint.app import live
    want_total = bool(re.search(r"\btotal\b", (live.request() or "").lower()))
    started = time.monotonic()
    label = note = ""
    if not path and (_URL.match(what.strip()) or what.strip().startswith(("~/", "/"))):
        path, what = what.strip(), ""                  # the model put the page or file in `what`
    try:
        if data.strip():
            # Data Mint already has (read from files or pages): straight in, no scratch .csv left behind.
            rows = _split_table(data)
            # A Total row the model wrote itself gets a formula or sum wrong (bench: =SUM over its own row
            # doubled the total). Take it out; a correct SUM row is added below.
            while rows and len(rows) > 1 and re.match(r"\s*(grand\s+)?totals?\b", str(rows[-1][0] or ""), re.I):
                rows.pop()
                want_total = True
            label = name or "the data"
            tables = [_table_from_rows(rows, "Data")] if rows else _extract_tables([data[:200_000]], what)
        elif _URL.match(path.strip()):
            label = path.strip()
            label, tables = _tables_from_url(label, what)
        elif path or source == "file":
            file, note = _file(path, SHEET_SOURCES)
            if file is None:
                return note
            label, tables = file.name, _tables_from_file(file, what)
        elif source == "clipboard":
            kind, value = _clipboard()
            if kind == "file":
                label, tables = value.name, _tables_from_file(value, what)
            elif kind == "image":
                label, tables = "the picture on the clipboard", _tables_from_image(value, what)
            elif kind == "text":
                rows = _split_table(value)
                label = "the copied text"
                tables = [_table_from_rows(rows, "Data")] if rows else _extract_tables([value[:200_000]], what)
            else:
                return "The clipboard is empty (or holds a password, which Mint does not read)."
        else:
            label = "the whole screen" if source == "screen" else "the front window"
            tables = _tables_from_image(_grab(source), what)
    except Exception as error:
        log.info("data_to_sheet: %s", error)
        return f"FAILED: could not get the data from {label or 'there'}: {error}"
    tables = [t for t in tables if t["rows"]]
    note = f" {note}" if note else ""
    if not tables:
        return f"Found no table or list of data in {label}.{note}"
    if job is not None:
        job["step"] = "writing the spreadsheet"
    fallback = Path(label).stem if "." in label else "Table"
    title = re.sub(r"[^\w .()-]", "", name or what or fallback).strip()[:60] or "Table"
    from mint.tools import saveto
    title = re.sub(r"\.(xlsx|xls|csv)$", "", title, flags=re.I).strip() or "Spreadsheet"
    out, moved = saveto.destination(config.storage("Spreadsheets") / f"{title}.xlsx", ".xlsx", save_to)
    if out.parent == config.storage("Spreadsheets"):
        out = config.storage("Spreadsheets") / f"{title} {time.strftime('%Y-%m-%d %H.%M.%S')}.xlsx"
    _write_book(out, tables)
    if want_total and data.strip():
        _add_total(out)
    if open_after:
        subprocess.run(["open", str(out)], check=False)
    sizes = "; ".join(f"'{t['name']}' {len(t['rows'])} rows x {max(len(t['columns']), 1)} columns "
                      f"({', '.join(str(c) for c in t['columns'][:6])}{'…' if len(t['columns']) > 6 else ''})"
                      for t in tables)
    return (f"Made a spreadsheet from {label} in {time.monotonic() - started:.0f} s: {sizes}. Saved: {out}"
            + (" and opened it." if open_after else ".") + note)


# ---------------------------------------------------------------- convert_document

def _body(page: str) -> tuple[str, str]:
    """(the <style> rules, the inside of <body>) of an HTML page."""
    styles = "\n".join(re.findall(r"<style[^>]*>(.*?)</style>", page, re.S | re.I))
    match = re.search(r"<body[^>]*>(.*)</body>", page, re.S | re.I)
    return styles, (match.group(1) if match else page).strip()


def _markdown(text: str) -> str:
    """Markdown (headings, lists, tables, bold, italics, code) as HTML."""
    def inline(line):
        line = html.escape(line)
        line = re.sub(r"`([^`]+)`", r"<code>\1</code>", line)
        line = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: f"<b>{m.group(1) or m.group(2)}</b>", line)
        line = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?!\w)", r"<i>\1</i>", line)
        return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', line)
    out, lines, i = [], text.splitlines(), 0
    paragraph: list[str] = []
    lists: list[tuple[str, int]] = []               # the open lists (kind, indent), each with an open item

    def flush():
        if paragraph:
            out.append("<p>" + inline(" ".join(paragraph)) + "</p>")
            paragraph.clear()
        while lists:
            out.append(f"</li></{lists.pop()[0]}>")
    while i < len(lines):
        line = lines[i].rstrip()
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        item = re.match(r"^\s*([-*+]|\d+[.)])\s+(.*)$", line)
        if line.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(html.escape(lines[i]))
                i += 1
            out.append("<pre>" + "\n".join(code) + "</pre>")
        elif heading:
            flush()
            level = len(heading.group(1))
            out.append(f"<h{level}>{inline(heading.group(2))}</h{level}>")
        elif line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?[\s:|-]+$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            i -= 1
            out.append("<table>" + "".join("<tr>" + "".join(f"<{t}>{inline(c)}</{t}>" for c in row) + "</tr>"
                                            for n, row in enumerate(rows) for t in ["th" if n == 0 else "td"])
                       + "</table>")
        elif item:
            if paragraph:
                out.append("<p>" + inline(" ".join(paragraph)) + "</p>")
                paragraph.clear()
            kind, indent = "ol" if item.group(1)[0].isdigit() else "ul", len(line) - len(line.lstrip())
            while lists and lists[-1][1] > indent:
                out.append(f"</li></{lists.pop()[0]}>")
            if lists and lists[-1][1] == indent:
                if lists[-1][0] == kind:
                    out.append("</li>")
                else:
                    out.append(f"</li></{lists.pop()[0]}>")
            if not lists or lists[-1][1] < indent:
                out.append(f"<{kind}>")
                lists.append((kind, indent))
            out.append(f"<li>{inline(item.group(2))}")
        elif not line.strip():
            flush()
        else:
            if lists:
                flush()
            paragraph.append(line.strip())
        i += 1
    flush()
    return "\n".join(out)


# A list item typed as text, as Word and textutil write them: a tab, a bullet or number, a tab.
_TYPED_BULLET = re.compile(r"^(?:\s|&nbsp;)*(?:[•●▪◦‣∙○■□➢►–-]|(\d{1,3})[.)])(?:\s|&nbsp;)+")


class _Cleaner(HTMLParser):
    """Any HTML - a web page, or what textutil makes of a Word or RTF file - as plain structure: one h1-h3, p,
    pre, list or table per line, with only b, i, u and br inside. In textutil's HTML a heading is a paragraph
    in a bigger font (its class) and a list item a paragraph starting with a tab and a bullet."""

    SKIP = {"script", "style", "head", "noscript", "svg", "template", "button", "select", "nav"}
    BLOCKS = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "pre", "dt", "dd", "section",
              "article", "header", "footer", "main", "aside", "figcaption", "caption", "address"}
    INLINE = {"b": "b", "strong": "b", "i": "i", "em": "i", "u": "u", "sup": "sup", "sub": "sub"}

    def __init__(self, styles: str = ""):
        super().__init__(convert_charrefs=True)
        self.sizes, self.bold, self.italic = {}, set(), set()
        for kind, name, rules in re.findall(r"(\w*)\.([\w-]+)\s*\{([^}]*)\}", styles):
            size = re.search(r"font(?:-size)?:[^;]*?([\d.]+)px", rules)
            if size:
                self.sizes[name] = float(size.group(1))
            if re.search(r"font-weight:\s*(bold|[6-9]00)|font:[^;]*\bbold\b|-Bold", rules):
                self.bold.add(name)
            if re.search(r"font-style:\s*italic|font:[^;]*\bitalic\b|-Italic|Oblique", rules):
                self.italic.add(name)
        self.blocks: list[dict] = []
        self.buffer: list[str] = []
        self.tag, self.size = "p", 0.0
        self.lists: list[str] = []
        self.closers: list[list[str]] = []           # what each open tag must close (a bold span -> </b>)
        self.skip = 0
        self.tables: list[list] = []                 # rows being read, innermost table last
        self.cell: list[str] | None = None
        self.cell_tag = "td"

    def _flush(self):
        text = "".join(self.buffer).strip()
        self.buffer = []
        if not re.sub(r"<[^>]+>|\s|&nbsp;", "", text):
            return
        tag, kind = self.tag, self.lists[-1] if self.lists else ""
        bullet = _TYPED_BULLET.match(text)
        if tag == "p" and bullet:                    # a Word list item
            tag, kind, text = "li", "ol" if bullet.group(1) else "ul", text[bullet.end():]
        self.blocks.append({"tag": tag, "html": re.sub(r"\s+", " ", text) if tag != "pre" else text,
                            "list": kind, "size": self.size, "depth": max(1, len(self.lists))})
        self.tag, self.size = "p", 0.0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
            return
        if self.skip:
            return
        attributes = dict(attrs)
        classes = (attributes.get("class") or "").split()
        if tag == "table":
            self._flush()
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self.tables[-1].append([])
        elif tag in ("td", "th") and self.tables:
            if len(self.tables) == 1:
                self.cell, self.cell_tag = [], tag
            if not self.tables[-1]:
                self.tables[-1].append([])
        elif tag in ("ul", "ol"):
            self._flush()
            self.lists.append(tag)
        elif tag in self.BLOCKS and self.cell is None:
            self._flush()
            if re.fullmatch(r"h[1-6]", tag):
                self.tag = f"h{min(3, int(tag[1]))}"
            elif tag in ("li", "pre"):
                self.tag = tag
            self.size = max((self.sizes.get(c, 0.0) for c in classes), default=0.0)
        elif tag == "br":
            self._write("<br>")
        if self.cell is not None and tag in ("li", "p", "div"):
            self._write(" ")
        closers = []
        if tag in self.INLINE:
            self._write(f"<{self.INLINE[tag]}>")
            closers.append(self.INLINE[tag])
        if any(c in self.bold for c in classes) and tag not in self.INLINE:
            self._write("<b>")
            closers.append("b")
        if any(c in self.italic for c in classes) and tag not in self.INLINE:
            self._write("<i>")
            closers.append("i")
        if tag not in ("br", "img", "hr", "meta", "link", "input", "col", "wbr"):
            self.closers.append(closers)

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip or tag in ("br", "img", "hr", "meta", "link", "input", "col", "wbr"):
            return
        for closer in reversed(self.closers.pop() if self.closers else []):
            self._write(f"</{closer}>")
        if tag in ("td", "th") and self.cell is not None and len(self.tables) == 1:
            if self.tables[-1]:
                self.tables[-1][-1].append((self.cell_tag, re.sub(r"\s+", " ", "".join(self.cell)).strip()))
            self.cell = None
        elif tag == "table" and self.tables:
            rows = [row for row in self.tables.pop() if row]
            if rows and not self.tables:
                self.blocks.append({"tag": "table", "rows": rows})
        elif tag in ("ul", "ol") and self.lists:
            self._flush()
            self.lists.pop()
        elif tag in self.BLOCKS and self.cell is None:
            self._flush()

    def _write(self, text):
        (self.cell if self.cell is not None else self.buffer).append(text)

    def handle_data(self, data):
        if not self.skip:
            self._write(html.escape(data if self.tag == "pre" else re.sub(r"\s+", " ", data), quote=False))

    def body(self) -> str:
        """The blocks as HTML, one per line; paragraphs in a bigger font become headings."""
        self._flush()
        sizes = {}
        for block in self.blocks:
            if block["tag"] == "p" and block["size"]:
                sizes[block["size"]] = sizes.get(block["size"], 0) + len(block["html"])
        body_size = max(sizes, key=sizes.get) if sizes else 0.0
        out, open_lists = [], []
        for block in self.blocks:
            tag = block["tag"]
            if tag == "p" and body_size and block["size"] >= 1.5 * body_size:
                tag = "h1"
            elif tag == "p" and body_size and block["size"] >= 1.2 * body_size:
                tag = "h2"
            elif tag == "p" and len(block.get("html", "")) < 120 and re.fullmatch(r"<b>[^<]*</b>", block["html"]):
                tag = "h3"                                              # a short bold line on its own
            depth = block.get("depth", 1) if tag == "li" else 0
            while len(open_lists) > depth or (open_lists and tag == "li" and open_lists[-1] != block["list"]
                                               and len(open_lists) == depth):
                out.append(f"</{open_lists.pop()}>")
            while tag == "li" and len(open_lists) < depth:
                open_lists.append(block["list"] or "ul")
                out.append(f"<{open_lists[-1]}>")
            if tag == "table":
                width = max(len(row) for row in block["rows"])
                out.append("<table>" + "".join(
                    "<tr>" + "".join(f"<{t}>{c}</{t}>" for t, c in row + [("td", "")] * (width - len(row))) + "</tr>"
                    for row in block["rows"]) + "</table>")
            else:
                text = block["html"]
                if tag in ("h1", "h2", "h3"):
                    text = re.sub(r"^<b>(.*)</b>$", r"\1", text)
                out.append(f"<{tag}>{text}</{tag}>")
        out.extend(f"</{kind}>" for kind in reversed(open_lists))
        return "\n".join(out)


def _clean(page: str) -> str:
    styles, _ = _body(page)
    cleaner = _Cleaner(styles)
    cleaner.feed(page)
    cleaner.close()
    return cleaner.body()


def _word_html(path: Path) -> str:
    """A .docx as clean HTML, read from its XML: textutil ignores Word's styles, so headings made with Heading 1-3
    (as most are) would come out as plain paragraphs. Lists come from Word's numbering, tables stay tables."""
    import zipfile
    from xml.etree import ElementTree
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())

        def part(name):
            return ElementTree.fromstring(package.read(name)) if name in names else None
        document, styles, numbering = part("word/document.xml"), part("word/styles.xml"), part("word/numbering.xml")
    if document is None:
        raise RuntimeError(f"{path.name} is not a Word document")
    levels: dict[str, int] = {}                      # style id -> heading level
    for style in styles.iter(f"{w}style") if styles is not None else []:
        name = (style.find(f"{w}name").get(f"{w}val", "") if style.find(f"{w}name") is not None else "").lower()
        outline = style.find(f"{w}pPr/{w}outlineLvl")
        heading = re.fullmatch(r"heading ?(\d)", name)
        level = (1 if name == "title" else 2 if name == "subtitle" else int(heading.group(1)) if heading
                 else int(outline.get(f"{w}val", "9")) + 1 if outline is not None else 0)
        if 0 < level <= 9:
            levels[style.get(f"{w}styleId", "")] = level
    ordered: set[tuple[str, str]] = set()            # (list id, level) that count 1, 2, 3 rather than bullets
    if numbering is not None:
        abstract = {a.get(f"{w}abstractNumId"): a for a in numbering.iter(f"{w}abstractNum")}
        for num in numbering.iter(f"{w}num"):
            link = num.find(f"{w}abstractNumId")
            for level in abstract.get(link.get(f"{w}val") if link is not None else "", []):
                shape = level.find(f"{w}numFmt") if level.tag == f"{w}lvl" else None
                if shape is not None and shape.get(f"{w}val") not in ("bullet", "none"):
                    ordered.add((num.get(f"{w}numId"), level.get(f"{w}ilvl")))

    def on(props, name):
        mark = props.find(f"{w}{name}") if props is not None else None
        return mark is not None and mark.get(f"{w}val", "true") not in ("0", "false", "none")

    def runs(element) -> str:
        out = []
        for child in element:
            if child.tag in (f"{w}del", f"{w}pPr", f"{w}rPr", f"{w}moveFrom"):
                continue
            if child.tag == f"{w}r":
                props, text = child.find(f"{w}rPr"), []
                for piece in child:
                    if piece.tag == f"{w}t":
                        text.append(html.escape(piece.text or "", quote=False))
                    elif piece.tag == f"{w}tab":
                        text.append(" ")
                    elif piece.tag in (f"{w}br", f"{w}cr"):
                        text.append("<br>")
                joined = "".join(text)
                for mark, tag in (("u", "u"), ("i", "i"), ("b", "b")):
                    if joined.strip() and on(props, mark):
                        joined = f"<{tag}>{joined}</{tag}>"
                out.append(joined)
            else:
                out.append(runs(child))            # hyperlinks, fields, insertions, content controls
        return re.sub(r"</(b|i|u)><\1>", "", "".join(out))

    blocks: list[dict] = []

    def walk(container):
        for child in container:
            if child.tag == f"{w}p":
                props = child.find(f"{w}pPr")
                style = props.find(f"{w}pStyle") if props is not None else None
                style_id = style.get(f"{w}val", "") if style is not None else ""
                number = props.find(f"{w}numPr") if props is not None else None
                outline = props.find(f"{w}outlineLvl") if props is not None else None
                level = levels.get(style_id) or (int(outline.get(f"{w}val", "9")) + 1 if outline is not None else 0)
                text = runs(child).strip()
                if not re.sub(r"<[^>]+>|\s", "", text):
                    continue
                if 0 < level <= 9:
                    blocks.append({"tag": f"h{min(3, level)}", "html": text, "list": "", "size": 0.0, "depth": 1})
                elif number is not None and number.find(f"{w}numId") is not None:
                    num_id = number.find(f"{w}numId").get(f"{w}val")
                    depth = number.find(f"{w}ilvl").get(f"{w}val", "0") if number.find(f"{w}ilvl") is not None \
                        else "0"
                    blocks.append({"tag": "li", "html": text, "size": 0.0, "depth": int(depth) + 1,
                                   "list": "ol" if (num_id, depth) in ordered else "ul"})
                elif _TYPED_BULLET.match(text):
                    bullet = _TYPED_BULLET.match(text)
                    blocks.append({"tag": "li", "html": text[bullet.end():], "size": 0.0, "depth": 1,
                                   "list": "ol" if bullet.group(1) else "ul"})
                else:
                    sizes = [int(size.get(f"{w}val", "0")) / 2 for size in child.iter(f"{w}sz")
                             if size.get(f"{w}val", "").isdigit()]
                    blocks.append({"tag": "p", "html": text, "list": "", "size": max(sizes, default=0.0),
                                   "depth": 1})
            elif child.tag == f"{w}tbl":
                rows = []
                for row in child.iter(f"{w}tr"):
                    header = on(row.find(f"{w}trPr"), "tblHeader")
                    cells = [("th" if header else "td",
                              "<br>".join(t for t in (runs(p).strip() for p in cell.iter(f"{w}p")) if t))
                             for cell in row.findall(f"{w}tc")]
                    if cells:
                        rows.append(cells)
                if rows:
                    blocks.append({"tag": "table", "rows": rows})
            elif child.tag in (f"{w}sdt", f"{w}sdtContent", f"{w}customXml", f"{w}ins"):
                walk(child)
    walk(document.find(f"{w}body"))
    cleaner = _Cleaner()
    cleaner.blocks = blocks
    return cleaner.body()


def _read_document(path: Path, job: dict) -> tuple[str, int, str]:
    """(body as clean HTML, pages, a note) for any document."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages, scanned = _pdf_blocks(path, job)
        body = "\n".join(_as_html(blocks) for blocks in pages)
        return body, len(pages), (f"{scanned} scanned page(s) read by text recognition" if scanned else "")
    if suffix in IMAGES:
        return _as_html(_blocks(_recognise(_image_file(path)))), 1, "read by text recognition"
    if suffix == ".docx":
        try:
            return _word_html(path), 0, ""
        except Exception as error:                     # an odd file: let textutil try
            log.info("reading %s: %s", path.name, error)
    if suffix in RICH:
        return _clean(_textutil(path, "html")), 0, ""
    text = path.read_text(errors="replace")
    if suffix in {".html", ".htm"}:
        return _clean(text), 0, ""
    if suffix in {".md", ".markdown"}:
        return _markdown(text), 0, ""
    blocks = ["<p>" + "<br>".join(html.escape(line, quote=False) for line in part.strip("\n").splitlines()) + "</p>"
              for part in re.split(r"\n\s*\n", text) if part.strip()]
    return "\n".join(blocks), 0, ""


def _pieces(body: str) -> list[str]:
    """The body in pieces of about CHUNK characters, cut only between elements."""
    parts: list[str] = []
    for line in re.split(r"\n+", body):
        while len(line) > CHUNK:                      # one long line: cut after a closing tag
            cut = max((m.end() for m in re.finditer(r"</(p|li|tr|h\d|div|ul|ol|table|blockquote)>|<br\s*/?>",
                                                     line[:CHUNK], re.I)), default=0)
            if cut == 0:
                cut = line.rfind(">", 0, CHUNK) + 1 or CHUNK
            parts.append(line[:cut])
            line = line[cut:]
        parts.append(line)
    chunks, current = [], ""
    for part in parts:
        if current and len(current) + len(part) > CHUNK:
            chunks.append(current)
            current = ""
        current += part + "\n"
    if current.strip():
        chunks.append(current)
    return chunks


def _texts(fragment: str) -> list[str]:
    """The pieces of real text between the tags (a word or more, not just numbers or marks)."""
    return [t.strip() for t in re.split(r"<[^>]+>", fragment) if re.search(r"[^\W\d_]{2,}", t)]


def _translate_piece(piece: str, language: str, other_script: bool = False) -> str:
    """One piece of HTML in another language. Tried again with the next model when tags went missing; text that
    came back untranslated (Flash Lite often leaves headings like "Section 5: District Progress" in English) is
    translated on its own afterwards. For a language in another script (Hindi, Tamil, Japanese), text still
    holding English words ("Section 9: जिला प्रगति") counts as left over too."""
    from mint.core import llm
    if not re.sub(r"<[^>]+>|&\w+;|\s|\d", "", piece):
        return piece                                   # nothing to translate (empty tags, numbers)
    prompt = (f"Translate this piece of an HTML document into {language}. Keep every tag and attribute exactly as "
              "it is, in the same order - translate only the text between the tags, all of it: headings, list "
              "items and table cells too. It may start or end in the middle of a list or table: do not add, close "
              "or remove any tags. Keep numbers, amounts and dates in digits, and names of people, brands, URLs, "
              f"email addresses and code as they are. Write natural, fluent {language} in its own script, as a "
              f"professional translator would, with the usual spaces between words. Text already in {language} "
              "stays. Return only the translated HTML - no explanation, no code fences.\n\n" + piece)
    tags = len(re.findall(r"<[a-zA-Z/]", piece))
    text = ""
    for attempt in range(2):
        answer, _model = llm.generate(prompt, MODELS[attempt:] if attempt else MODELS)
        answer = re.sub(r"^\s*```(?:html)?\s*|\s*```\s*$", "", answer.strip())
        if len(re.findall(r"<[a-zA-Z/]", answer)) >= 0.8 * tags and len(answer) > 0.2 * len(piece):
            text = answer
            break
        text = text or answer
        log.info("translation try %d lost tags (%d chars)", attempt + 1, len(piece))
    if not text:
        return piece
    source = set(_texts(piece))
    left = sorted({t for t in _texts(text) if t in source or (other_script and re.search(r"[A-Za-z]{3,}", t))},
                  key=len, reverse=True)
    if left:
        log.info("%d of %d texts came back untranslated; translating them on their own", len(left), len(source))
        try:
            answer, _model = llm.generate(
                f"Translate each of these texts from a document into natural {language}, in its own script. Keep "
                "numbers, names, URLs and any HTML entities as they are. Return JSON {\"translations\": [...]} with "
                "exactly one translation per text, in the same order.\n\n" + json.dumps(left, ensure_ascii=False),
                MODELS, json_mode=True)
            done = llm.parse_json(answer).get("translations") or []
            if len(done) == len(left):
                for original, translated in zip(left, done):
                    if isinstance(translated, str) and translated.strip():
                        text = re.sub(rf"(^|>)(\s*){re.escape(original)}(\s*)(?=<|$)",       # whole texts only
                                      lambda m, new=translated.strip(): m.group(1) + m.group(2) + new + m.group(3),
                                      text, flags=re.M)
        except Exception as error:
            log.info("translating the leftovers: %s", error)
    return text


def _language(name: str) -> tuple[str, str, str, bool]:
    """(language name, code, font, right to left) for what the user said ("hindi", "hi", "Hindi")."""
    from mint.tools.translate import LANGUAGES, _language as full_name
    label = full_name(name) if name.strip().lower() in LANGUAGES else name.strip().title()
    code, font, rtl = SCRIPTS.get(label.lower(), (next((k for k, v in LANGUAGES.items() if v == label), ""),
                                                  "", False))
    return label, code, font, rtl


def _page(title: str, body: str, code: str, font: str, rtl: bool) -> str:
    """A whole HTML page: UTF-8, the language marked, and a font that has its script."""
    family = f"'{font}', '{LATIN_FONT}', sans-serif" if font else f"'{LATIN_FONT}', Helvetica, Arial, sans-serif"
    style = (f"body {{ font-family: {family}; font-size: 11.5pt; line-height: 1.5; margin: 0; }}\n"
             "h1 { font-size: 20pt; } h2 { font-size: 15pt; } h3 { font-size: 12.5pt; }\n"
             "table { border-collapse: collapse; margin: 6pt 0 10pt; }\n"
             "th, td { border: 1px solid #999; padding: 3pt 6pt; vertical-align: top; text-align: start; }\n"
             "th { background: #eee; } pre { font: 9.5pt Menlo, monospace; white-space: pre-wrap; }\n"
             "@page { margin: 20mm 18mm; }")
    direction = ' dir="rtl"' if rtl else ""
    return (f'<!doctype html>\n<html lang="{code or "en"}"{direction}>\n<head>\n<meta charset="utf-8">\n'
            f"<title>{html.escape(title)}</title>\n<style>\n{style}\n</style>\n</head>\n"
            f"<body>\n{body}\n</body>\n</html>\n")


class _ToMarkdown(HTMLParser):
    """HTML -> Markdown for headings, paragraphs, lists, tables, bold and italics; or, plain, -> text with
    bullets and tab-separated table cells."""

    def __init__(self, plain: bool = False):
        super().__init__(convert_charrefs=True)
        self.plain = plain
        self.out: list[str] = []
        self.lists: list[str] = []
        self.row: list[str] | None = None
        self.rows: list[list[str]] = []
        self.cell: list[str] | None = None
        self.skip = 0

    def _write(self, text):
        (self.cell if self.cell is not None else self.out).append(text)

    def handle_starttag(self, tag, attrs):
        if tag in ("style", "script", "head", "title"):
            self.skip += 1
        elif re.fullmatch(r"h[1-6]", tag):
            self._write("\n\n" + ("" if self.plain else "#" * int(tag[1]) + " "))
        elif tag in ("p", "div"):
            self._write("\n\n")
        elif tag == "br":
            self._write("  \n")
        elif tag in ("ul", "ol"):
            self.lists.append(tag)
            self._write("\n")
        elif tag == "li":
            depth = max(0, len(self.lists) - 1)
            self._write("\n" + "  " * depth + ("1. " if self.lists[-1:] == ["ol"] else "• " if self.plain else "- "))
        elif tag in ("b", "strong") and not self.plain:
            self._write("**")
        elif tag in ("i", "em") and not self.plain:
            self._write("*")
        elif tag == "tr":
            self.row = []
        elif tag in ("td", "th"):
            self.cell = []

    def handle_endtag(self, tag):
        if tag in ("style", "script", "head", "title"):
            self.skip = max(0, self.skip - 1)
        elif tag in ("ul", "ol") and self.lists:
            self.lists.pop()
            self._write("\n")
        elif tag in ("b", "strong") and not self.plain:
            self._write("**")
        elif tag in ("i", "em") and not self.plain:
            self._write("*")
        elif tag in ("td", "th") and self.cell is not None and self.row is not None:
            cell = " ".join("".join(self.cell).split())
            self.row.append(cell if self.plain else cell.replace("|", "\\|"))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None
        elif tag == "table" and self.rows:
            width = max(len(r) for r in self.rows)
            if self.plain:
                lines = ["\t".join(r) for r in self.rows]
            else:
                lines = ["| " + " | ".join(r + [""] * (width - len(r))) + " |" for r in self.rows]
                lines.insert(1, "|" + " --- |" * width)
            self.out.append("\n\n" + "\n".join(lines) + "\n\n")
            self.rows = []

    def handle_data(self, data):
        if not self.skip:
            self._write(re.sub(r"\s+", " ", data))

    def text(self) -> str:
        text = "".join(self.out)
        if not self.plain:
            text = re.sub(r"\*\*\s*\*\*|(?<![*\w])\*\s*\*(?![*\w])", "", text)       # empty bold or italics
        text = re.sub(r"[ \t]+\n", "\n", text)
        return re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"


def _pdf(page_path: Path, target: Path) -> None:
    """HTML -> PDF: headless Chrome in a profile of its own (draws every script well), else the Mac's own text
    system in a separate process."""
    from mint.tools.documents import CHROME
    target.unlink(missing_ok=True)
    work = Path(tempfile.mkdtemp(prefix="mint-convert-"))
    try:
        if Path(CHROME).exists():
            process = subprocess.Popen(
                [CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--no-first-run",
                 f"--user-data-dir={work / 'profile'}", f"--print-to-pdf={target}", page_path.as_uri()],
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            # Headless Chrome writes the file and then often does not exit: wait for the file to settle.
            deadline, last = time.monotonic() + 120, -1
            while time.monotonic() < deadline:
                time.sleep(0.4)
                size = target.stat().st_size if target.exists() else -1
                if size > 0 and size == last:
                    break
                last = size
            process.terminate()
            try:
                process.wait(timeout=3)
            except subprocess.TimeoutExpired:
                process.kill()
        if not target.exists() or target.stat().st_size == 0:
            import sys
            subprocess.run([sys.executable, "-c", _APPKIT_PDF, str(page_path), str(target)], capture_output=True,
                           timeout=180, check=False)
    finally:
        shutil.rmtree(work, ignore_errors=True)
    if not target.exists() or target.stat().st_size == 0:
        raise RuntimeError("the PDF could not be made")


_APPKIT_PDF = """
import sys, AppKit, Foundation
data = Foundation.NSData.dataWithContentsOfFile_(sys.argv[1])
text, _ = AppKit.NSAttributedString.alloc().initWithHTML_documentAttributes_(data, None)
info = AppKit.NSPrintInfo.sharedPrintInfo().copy()
info.setJobDisposition_(AppKit.NSPrintSaveJob)
info.dictionary()[AppKit.NSPrintJobSavingURL] = Foundation.NSURL.fileURLWithPath_(sys.argv[2])
for side in ("Left", "Right", "Top", "Bottom"):
    getattr(info, "set%sMargin_" % side)(54)
info.setHorizontalPagination_(AppKit.NSFitPagination)
size = info.paperSize()
view = AppKit.NSTextView.alloc().initWithFrame_(((0, 0), (size.width - 108, size.height - 108)))
view.textStorage().setAttributedString_(text)
view.sizeToFit()
operation = AppKit.NSPrintOperation.printOperationWithView_printInfo_(view, info)
operation.setShowsPrintPanel_(False)
operation.setShowsProgressPanel_(False)
operation.runOperation()
"""


class _WordWriter(HTMLParser):
    """Clean HTML (see _Cleaner) -> the body of a Word document (WordprocessingML)."""

    def __init__(self, rtl: bool):
        super().__init__(convert_charrefs=True)
        self.rtl = rtl
        self.out: list[str] = []
        self.runs: list[str] = []
        self.style = ""
        self.marks = {"b": 0, "i": 0, "u": 0, "sup": 0, "sub": 0}
        self.lists: list[int] = []                   # numbering ids of the open lists
        self.numbers = 2                             # 1 = bullets, 2 = the first numbered list
        self.ordered: list[int] = []                 # numbered lists, each restarting at 1
        self.rows: list[list[tuple[bool, list[str]]]] | None = None
        self.cell: list[str] | None = None
        self.header = False
        self.pre = False

    def _paragraph(self):
        if not self.runs:
            return
        props = f'<w:pStyle w:val="{self.style}"/>' if self.style else ""
        if self.style == "ListParagraph" and self.lists:
            props += f'<w:numPr><w:ilvl w:val="{min(2, len(self.lists) - 1)}"/><w:numId w:val="{self.lists[-1]}"/></w:numPr>'
        if self.rtl:
            props += "<w:bidi/>"
        xml = f"<w:p><w:pPr>{props}</w:pPr>{''.join(self.runs)}</w:p>"
        (self.cell if self.cell is not None else self.out).append(xml)
        self.runs = []

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "h3", "p", "li", "pre", "div", "blockquote"):
            self._paragraph()
            self.style = {"h1": "Heading1", "h2": "Heading2", "h3": "Heading3", "li": "ListParagraph",
                          "pre": "Code"}.get(tag, "")
            self.pre = tag == "pre"
        elif tag in ("ul", "ol"):
            self._paragraph()
            if tag == "ol":
                self.numbers += 1
                self.ordered.append(self.numbers)
            self.lists.append(self.numbers if tag == "ol" else 1)
        elif tag == "table":
            self._paragraph()
            self.rows = []
        elif tag == "tr" and self.rows is not None:
            self.rows.append([])
        elif tag in ("td", "th") and self.rows is not None:
            self.cell, self.header = [], tag == "th"
            if not self.rows:
                self.rows.append([])
        elif tag == "br":
            self.runs.append("<w:r><w:br/></w:r>")
        elif tag in self.marks:
            self.marks[tag] += 1

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "h3", "p", "li", "pre", "div", "blockquote"):
            self._paragraph()
            self.style, self.pre = "", False
        elif tag in ("ul", "ol") and self.lists:
            self._paragraph()
            self.lists.pop()
        elif tag in ("td", "th") and self.cell is not None and self.rows is not None:
            self._paragraph()
            self.rows[-1].append((self.header, self.cell or ["<w:p/>"]))
            self.cell, self.header = None, False
        elif tag == "table" and self.rows is not None:
            self._table([row for row in self.rows if row])
            self.rows = None
        elif tag in self.marks:
            self.marks[tag] = max(0, self.marks[tag] - 1)

    def handle_data(self, data):
        text = data if self.pre else re.sub(r"\s+", " ", data)
        if not text or (not text.strip() and not self.runs):
            return
        marks = ("<w:b/><w:bCs/>" if self.marks["b"] or self.header else "") + \
                ("<w:i/><w:iCs/>" if self.marks["i"] else "") + ('<w:u w:val="single"/>' if self.marks["u"] else "") + \
                ('<w:vertAlign w:val="superscript"/>' if self.marks["sup"] else "") + \
                ('<w:vertAlign w:val="subscript"/>' if self.marks["sub"] else "") + ("<w:rtl/>" if self.rtl else "")
        for n, line in enumerate(text.split("\n")):
            if n:
                self.runs.append("<w:r><w:br/></w:r>")
            self.runs.append(f'<w:r><w:rPr>{marks}</w:rPr><w:t xml:space="preserve">{html.escape(line, quote=False)}'
                             "</w:t></w:r>")

    def _table(self, rows):
        if not rows:
            return
        width = max(len(row) for row in rows)
        column = 9600 // width
        border = "".join(f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="999999"/>'
                         for side in ("top", "left", "bottom", "right", "insideH", "insideV"))
        xml = [f'<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/>{"<w:bidiVisual/>" if self.rtl else ""}'
               f'<w:tblBorders>{border}</w:tblBorders><w:tblCellMar><w:left w:w="80" w:type="dxa"/>'
               '<w:right w:w="80" w:type="dxa"/></w:tblCellMar></w:tblPr><w:tblGrid>'
               + f'<w:gridCol w:w="{column}"/>' * width + "</w:tblGrid>"]
        for row in rows:
            header = all(is_header for is_header, _ in row)
            xml.append("<w:tr>" + ("<w:trPr><w:tblHeader/></w:trPr>" if header else ""))
            for is_header, paragraphs in row + [(False, ["<w:p/>"])] * (width - len(row)):
                shade = '<w:shd w:val="clear" w:color="auto" w:fill="EEEEEE"/>' if is_header else ""
                xml.append(f'<w:tc><w:tcPr><w:tcW w:w="{column}" w:type="dxa"/>{shade}</w:tcPr>'
                           + "".join(paragraphs) + "</w:tc>")
            xml.append("</w:tr>")
        xml.append("</w:tbl><w:p/>")
        self.out.append("".join(xml))

    def document(self) -> str:
        self._paragraph()
        return "".join(self.out)


_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _docx(body: str, target: Path, title: str, code: str, font: str, rtl: bool) -> None:
    """Clean HTML -> a Word document, written here because textutil's Word export drops tables: Heading 1-3,
    bullet and numbered lists, bordered tables, bold, italics, underline, the language marked and a font for its
    script (Word picks the "complex script" font for Indian and Arabic scripts, so it is set for all four)."""
    import zipfile
    writer = _WordWriter(rtl)
    writer.feed(body)
    writer.close()
    face = font or LATIN_FONT
    lang = code or "en-US"
    fonts = f'<w:rFonts w:ascii="{face}" w:hAnsi="{face}" w:cs="{face}" w:eastAsia="{face}"/>'
    languages = f'<w:lang w:val="{lang}" w:eastAsia="{lang}" w:bidi="{lang}"/>'

    def heading(level, size):
        return (f'<w:style w:type="paragraph" w:styleId="Heading{level}"><w:name w:val="heading {level}"/>'
                '<w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/>'
                f'<w:spacing w:before="{360 - level * 60}" w:after="120"/><w:outlineLvl w:val="{level - 1}"/></w:pPr>'
                f'<w:rPr><w:b/><w:bCs/><w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr></w:style>')
    styles = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles {_W}><w:docDefaults><w:rPrDefault>'
              f'<w:rPr>{fonts}<w:sz w:val="22"/><w:szCs w:val="22"/>{languages}</w:rPr></w:rPrDefault>'
              '<w:pPrDefault><w:pPr><w:spacing w:after="140" w:line="288" w:lineRule="auto"/></w:pPr></w:pPrDefault>'
              '</w:docDefaults><w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/>'
              '<w:qFormat/></w:style>' + heading(1, 36) + heading(2, 30) + heading(3, 25)
              + '<w:style w:type="paragraph" w:styleId="ListParagraph"><w:name w:val="List Paragraph"/>'
              '<w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="60"/><w:contextualSpacing/></w:pPr></w:style>'
              '<w:style w:type="paragraph" w:styleId="Code"><w:name w:val="Code"/><w:basedOn w:val="Normal"/>'
              '<w:pPr><w:spacing w:after="0"/></w:pPr><w:rPr><w:rFonts w:ascii="Menlo" w:hAnsi="Menlo"/>'
              '<w:sz w:val="19"/></w:rPr></w:style></w:styles>')

    def levels(ordered):
        return "".join(
            f'<w:lvl w:ilvl="{n}"><w:start w:val="1"/><w:numFmt w:val="{"decimal" if ordered else "bullet"}"/>'
            f'<w:lvlText w:val="{f"%{n + 1}." if ordered else "•◦▪"[n]}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{720 * (n + 1)}" w:hanging="360"/></w:pPr></w:lvl>' for n in range(3))
    numbering = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:numbering {_W}>'
                 f'<w:abstractNum w:abstractNumId="0">{levels(False)}</w:abstractNum>'
                 f'<w:abstractNum w:abstractNumId="1">{levels(True)}</w:abstractNum>'
                 '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
                 + "".join(f'<w:num w:numId="{n}"><w:abstractNumId w:val="1"/><w:lvlOverride w:ilvl="0">'
                           '<w:startOverride w:val="1"/></w:lvlOverride></w:num>' for n in writer.ordered)
                 + "</w:numbering>")
    document = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {_W}><w:body>'
                + writer.document()
                + '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1134" w:right="1134" w:bottom="1134" '
                'w:left="1134" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr></w:body></w:document>')
    relation = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package = "http://schemas.openxmlformats.org/package/2006"
    main = "application/vnd.openxmlformats-officedocument.wordprocessingml"
    files = {
        "[Content_Types].xml": (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="{package}/content-types">'
            f'<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/word/document.xml" ContentType="{main}.document.main+xml"/>'
            f'<Override PartName="/word/styles.xml" ContentType="{main}.styles+xml"/>'
            f'<Override PartName="/word/numbering.xml" ContentType="{main}.numbering+xml"/>'
            '<Override PartName="/docProps/core.xml" '
            'ContentType="application/vnd.openxmlformats-package.core-properties+xml"/></Types>'),
        "_rels/.rels": (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{package}/relationships">'
            f'<Relationship Id="rId1" Type="{relation}/officeDocument" Target="word/document.xml"/>'
            f'<Relationship Id="rId2" Type="{package}/relationships/metadata/core-properties" '
            'Target="docProps/core.xml"/></Relationships>'),
        "docProps/core.xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><cp:coreProperties xmlns:cp="http://schemas.'
            'openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/">'
            f"<dc:title>{html.escape(title, quote=False)}</dc:title><dc:language>{lang}</dc:language>"
            "</cp:coreProperties>"),
        "word/_rels/document.xml.rels": (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{package}/relationships">'
            f'<Relationship Id="rId1" Type="{relation}/styles" Target="styles.xml"/>'
            f'<Relationship Id="rId2" Type="{relation}/numbering" Target="numbering.xml"/></Relationships>'),
        "word/document.xml": document, "word/styles.xml": styles, "word/numbering.xml": numbering}
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as package_file:
        for name, text in files.items():
            package_file.writestr(name, text.encode("utf-8"))


def _write_output(body: str, target: Path, fmt: str, title: str, code: str, font: str, rtl: bool) -> None:
    page = _page(title, body, code, font, rtl)
    if fmt == "html":
        target.write_text(page, encoding="utf-8")
    elif fmt in ("md", "txt"):
        parser = _ToMarkdown(plain=fmt == "txt")
        parser.feed(body)
        target.write_text(parser.text(), encoding="utf-8")
    elif fmt == "docx":
        _docx(body, target, title, code, font, rtl)
    else:
        work = Path(tempfile.mkdtemp(prefix="mint-convert-"))
        try:
            source = work / "document.html"
            source.write_text(page, encoding="utf-8")
            if fmt == "pdf":
                _pdf(source, target)
            else:
                done = subprocess.run(["textutil", "-convert", fmt, "-inputencoding", "UTF-8", "-output",
                                       str(target), str(source)], capture_output=True, timeout=300, check=False)
                if done.returncode != 0 or not target.exists():
                    raise RuntimeError(f"textutil could not write {fmt}: {done.stderr.decode(errors='replace')[:200]}")
        finally:
            shutil.rmtree(work, ignore_errors=True)


def _target(path: Path, label: str, fmt: str) -> Path:
    """Next to the original ("report (Hindi).docx"), or in Mint's Documents folder when that folder is not
    writable. Never over an existing file."""
    from mint.core import config
    folder = path.parent if os.access(path.parent, os.W_OK) else config.storage("Documents")
    stem = f"{path.stem} ({label})" if label else path.stem
    target, n = folder / f"{stem}.{fmt}", 2
    while target.exists():
        target, n = folder / f"{stem} {n}.{fmt}", n + 1
    return target


def convert_document(path: str, language: str = "", fmt: str = "docx", open_after: bool = True,
                     job: dict | None = None) -> str:
    started = time.monotonic()
    job = job if job is not None else {}
    fmt = (fmt or "docx").lower().lstrip(".").replace("word", "docx").replace("markdown", "md")
    fmt = {"doc": "docx", "text": "txt", "htm": "html"}.get(fmt, fmt)
    if fmt not in FORMATS:
        return f"FAILED: cannot write '{fmt}'; choose one of {', '.join(FORMATS)}."
    file, why = _file(path)
    if file is None:
        return why
    suffix = file.suffix.lower()
    if suffix not in {".pdf"} | RICH | PLAIN | set(IMAGES):
        return f"FAILED: {file.name} is not a document Mint can convert (PDF, Word, RTF, text, Markdown, HTML, image)."
    label, code, font, rtl = _language(language) if language.strip() else ("", "", "", False)
    if not label and suffix.lstrip(".") == fmt:
        return f"{file.name} is already a {fmt} file. Say a language to translate it into, or another format."
    job.update(title=file.name, step="reading", done=0, total=0)
    try:
        body, pages, note = _read_document(file, job)
    except Exception as error:
        log.info("convert %s: %s", file.name, error)
        return f"FAILED: could not read {file.name}: {error}"
    if not re.sub(r"<[^>]+>|\s", "", body):
        return f"FAILED: found no text in {file.name}."
    if label:
        pieces = _pieces(body)
        job.update(step=f"translating into {label}", done=0, total=len(pieces))

        failed: list[str] = []

        def translate(piece):
            """A piece, waiting out Gemini's busy spells (per-minute quota, 503s) for up to about two minutes; a
            piece that still fails stays in the original language rather than losing the whole document."""
            try:
                for attempt in range(5):
                    try:
                        return _translate_piece(piece, label, bool(font))
                    except Exception as error:
                        if attempt == 4 or "PerDay" in str(error):
                            raise
                        job["step"] = f"translating into {label} (Gemini is busy, waiting)"
                        time.sleep(30)
            except Exception as error:
                log.info("translate %s: %s", file.name, error)
                failed.append(str(error)[:160])
                return piece
            finally:
                job["done"] = job.get("done", 0) + 1
                job["step"] = f"translating into {label}"
        with ThreadPoolExecutor(3) as pool:
            body = "\n".join(pool.map(translate, pieces))
        if len(failed) == len(pieces):
            return f"FAILED: Gemini could not translate {file.name}: {failed[0]}"
        if failed:
            note = (note + "; " if note else "") + (f"{len(failed)} of {len(pieces)} parts could not be translated "
                                                    f"and are in the original language ({failed[0][:80]})")
    job.update(step=f"writing {fmt}")
    target = _target(file, label, fmt)
    try:
        _write_output(body, target, fmt, f"{file.stem} ({label})" if label else file.stem, code, font, rtl)
    except Exception as error:
        return f"FAILED: could not write the {fmt} file: {error}"
    if open_after:
        subprocess.run(["open", str(target)], check=False)
    size = f"{pages} page(s)" if pages else f"{len(re.sub(r'<[^>]+>', '', body)):,} characters"
    return (f"Converted {file.name} ({size}" + (f"; {note}" if note else "") + ") "
            + (f"into {label}, " if label else "") + f"as {fmt.upper()} in {time.monotonic() - started:.0f} s. "
            f"Saved: {target}" + (" and opened it." if open_after else "."))


# ---------------------------------------------------------------- running, waiting, reporting

def snapshot() -> dict:
    """For the island: the job running now - {"kind", "title", "step", "done", "total", "seconds"} - else {}."""
    with _jobs_lock:
        running = [job for job in _jobs.values() if job.get("running")]
    if not running:
        return {}
    job = running[0]
    return {"kind": job["kind"], "title": job.get("title", ""), "step": job.get("step", ""),
            "done": job.get("done", 0), "total": job.get("total", 0),
            "seconds": time.monotonic() - job["started"]}


def _status() -> str:
    with _jobs_lock:
        jobs = list(_jobs.values())
    if not jobs:
        return "No conversion is running."
    lines = []
    for job in jobs:
        if job.get("running"):
            count = f" ({job['done']} of {job['total']})" if job.get("total") else ""
            lines.append(f"{job.get('title') or job['kind']}: {job.get('step', 'working')}{count}, "
                         f"{time.monotonic() - job['started']:.0f} s so far.")
        else:
            lines.append(f"Finished: {job.get('result', '')[:300]}")
    return "\n".join(lines)


def _run(kind: str, title: str, work, later: str) -> str:
    """Answer within WAIT seconds; otherwise say so, finish in the background and tell Mint when done."""
    key = f"{kind}-{time.monotonic():.3f}"
    job = {"kind": kind, "title": title, "step": "starting", "started": time.monotonic(), "running": True}
    state: dict = {}
    lock = threading.Lock()
    with _jobs_lock:
        for old in [k for k, v in _jobs.items() if not v.get("running")]:
            del _jobs[old]
        _jobs[key] = job

    def run():
        try:
            result = work(job)
        except Exception as error:
            log.exception("%s failed", kind)
            result = f"FAILED: {error}"
        job.update(running=False, result=result)
        with lock:
            state["result"] = result
            late = state.get("late")
        if late:
            try:
                from mint.tools.work import _notify_mint
                _notify_mint(f"(Mint's {kind.replace('_', ' ')} job, not the user.) {result} Tell the user briefly.")
            except Exception:
                pass
    worker = threading.Thread(target=run, daemon=True, name=kind)
    worker.start()
    worker.join(WAIT)
    with lock:
        if "result" in state:
            return state["result"]
        state["late"] = True
    count = f" ({job['done']} of {job['total']} {'pages' if 'page' in job.get('step', '') else 'parts'} so far)" \
        if job.get("total") else ""
    return f"{later}{count} It continues in the background; a message comes when it is done. Tell the user briefly."


def _ocr_tool(args: dict) -> str:
    source = str(args.get("source") or "window")
    path = str(args.get("path") or "")
    return _run("ocr_copy", path or source,
                lambda job: ocr_copy(source, path, bool(args.get("lines")), job, str(args.get("save_to") or "")),
                "Still reading the text.")


def _sheet_tool(args: dict) -> str:
    source = str(args.get("source") or "window")
    path = str(args.get("path") or "")
    from mint.tools import saveto
    # Worked out now, while the request is at hand (the job may finish after the next request).
    save_to = str(args.get("save_to") or "") or saveto.requested(".xlsx")
    return _run("data_to_sheet", path or source,
                lambda job: data_to_sheet(source, path, str(args.get("what") or ""), str(args.get("name") or ""),
                                          args.get("open") is not False, job, save_to, str(args.get("data") or "")),
                "Still reading the data for the spreadsheet.")


def _convert_tool(args: dict) -> str:
    if args.get("status"):
        return _status()
    path = str(args.get("path") or "")
    language = str(args.get("language") or "")
    fmt = str(args.get("format") or "docx")
    what = (f"into {language.strip().title()} " if language.strip() else "") + f"as {fmt.upper()}"
    return _run("convert_document", Path(path).name or "document",
                lambda job: convert_document(path, language, fmt, args.get("open") is not False, job),
                f"Converting {Path(path).name or 'the document'} {what}; a long document takes a few minutes.")


PROMPT = """Text and documents: "copy the text on my screen / from this image or PDF", "OCR this" -> ocr_copy \
(source window, screen, clipboard_image, or path; lines=true keeps the exact line breaks; "and save it to X.txt" \
-> save_to, so the whole text is saved, not what you retype). "Make an Excel of this \
table / this data / from this PDF" -> data_to_sheet (source window, screen, clipboard, or path). A file or web \
page the user names goes in path (a web page as its URL - it is fetched, no need to open it); source window only \
for what is in the front window now. "Convert this PDF to Word", "translate report.docx into Hindi as a PDF" -> \
convert_document with path (or "this" for the file selected in Finder or open in front), language only if they \
want it translated, and format (default docx); "how far is the conversion?" -> convert_document status=true."""


def declarations():
    from google.genai import types
    S, B = types.Type.STRING, types.Type.BOOLEAN
    path_note = ("a file path or name ('~/Downloads/report.pdf', 'downloads/scan.png'), or 'this' for the file "
                 "selected in Finder or open in the front app")
    return [
        types.FunctionDeclaration(
            name="ocr_copy",
            description=("Read the text of the front window, the whole screen, an image or PDF (scans too), or a "
                         "picture on the clipboard, with the Mac's own text recognition (any language), and copy it "
                         "to the clipboard in reading order, paragraphs kept."),
            parameters=types.Schema(type=types.Type.OBJECT, properties={
                "source": types.Schema(type=S, enum=["window", "screen", "file", "clipboard_image"],
                                       description="where the text is (default window; file when path is given)"),
                "path": types.Schema(type=S, description=path_note),
                "lines": types.Schema(type=B, description="keep the exact line breaks instead of joining "
                                                         "paragraphs (default false)"),
                "save_to": types.Schema(type=S, description="also save the text in this new text file "
                                                           "('~/Documents/poster.txt'), when the user wants it "
                                                           "saved")})),
        types.FunctionDeclaration(
            name="data_to_sheet",
            description=("Turn a table or list of data into an Excel .xlsx (numbers as numbers, dates as dates, a "
                         "bold frozen header, one sheet per table) and open it: from the front window or screen, a "
                         "PDF, image, Word, text or CSV file, a web page (its URL in path; fetched here, it does "
                         "not need to be open), or what is on the clipboard (copied cells, text or a picture). A "
                         "file or page the user names always goes in path."),
            parameters=types.Schema(type=types.Type.OBJECT, properties={
                "source": types.Schema(type=S, enum=["window", "screen", "file", "clipboard"],
                                       description="where the data is (default window; file when path is given)"),
                "path": types.Schema(type=S, description=path_note + ", or a web page URL ('https://...')"),
                "what": types.Schema(type=S, description="which table or data, if the user said, e.g. "
                                                         "'the price list', 'only the transactions'"),
                "name": types.Schema(type=S, description="file name for the spreadsheet (optional)"),
                "save_to": types.Schema(type=S, description="where the user asked for it: a folder or a full "
                                                            ".xlsx path; empty = Mint's Spreadsheets folder"),
                "data": types.Schema(type=S, description="the table itself as CSV/TSV text, when you already have "
                                                         "the data (from files or pages you read) - no scratch "
                                                         ".csv file needed"),
                "open": types.Schema(type=B, description="open it when done (default true)")})),
        types.FunctionDeclaration(
            name="convert_document",
            description=("Convert a document (PDF, scanned PDF, Word, RTF, text, Markdown, HTML, image) into Word "
                         "(docx), PDF, Markdown, text, HTML or RTF, keeping headings, lists, tables and paragraphs; "
                         "optionally translate it into another language (Hindi and other scripts included). Saved "
                         "next to the original, e.g. 'report (Hindi).docx', and opened. Long documents finish in "
                         "the background with a message."),
            parameters=types.Schema(type=types.Type.OBJECT, properties={
                "path": types.Schema(type=S, description=path_note),
                "language": types.Schema(type=S, description="translate into this language (only if asked)"),
                "format": types.Schema(type=S, enum=list(FORMATS), description="output format (default docx)"),
                "open": types.Schema(type=B, description="open it when done (default true)"),
                "status": types.Schema(type=B, description="only report how far the running conversion is")})),
    ]


HANDLERS = {"ocr_copy": _ocr_tool, "data_to_sheet": _sheet_tool, "convert_document": _convert_tool}
