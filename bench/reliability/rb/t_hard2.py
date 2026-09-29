"""Harder tasks (ids "hard-"): long chains, multi-turn conversations, messy data, constraints, recovery,
undo, agent handoff and clipboard + files. Every check reads the end state only."""

from __future__ import annotations

import csv
import datetime as dt
import io
import os
import re
import stat
import subprocess
from pathlib import Path

from . import apple as ap
from . import sandbox as sb
from .registry import Ctx, check, fail, ok, setup
from .t_files import image_bytes, pdf_bytes
from .t_sheets import number

# --- helpers ------------------------------------------------------------------------------------------


def _ago(days: float) -> dt.datetime:
    return dt.datetime.now() - dt.timedelta(days=days)


def _lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines() if line.strip()]


def _bullet(line: str) -> str:
    """A line without a list marker ("- ", "* ", "1. ", "[ ] ")."""
    return re.sub(r"^\s*(?:[-*+•]|\d+[.)])\s+(?:\[[ xX]\]\s+)?", "", line).strip()


def _text(path: Path) -> str | None:
    return path.read_text(errors="replace") if path.exists() else None


def _csv(path: Path) -> list[list[str]]:
    rows = list(csv.reader(io.StringIO(path.read_text(errors="replace"))))
    return [[c.strip() for c in r] for r in rows if any(c.strip() for c in r)]


def _money(text: str) -> float:
    return float(text.replace(",", "")) if text else 0.0


def _numbers(text: str) -> list[float]:
    return [float(n.replace(",", "")) for n in re.findall(r"(?<![\d.])\d[\d,]*(?:\.\d+)?", text)]


def _pdf(path: Path, lines: list[str]) -> Path:
    """A one-page PDF with a real text layer (CoreGraphics + CoreText)."""
    import CoreText
    import Quartz
    from Foundation import NSAttributedString, NSURL
    path.parent.mkdir(parents=True, exist_ok=True)
    context = Quartz.CGPDFContextCreateWithURL(NSURL.fileURLWithPath_(str(path)), ((0, 0), (612, 792)), None)
    Quartz.CGPDFContextBeginPage(context, None)
    font = CoreText.CTFontCreateWithName("Helvetica", 13, None)
    y = 740
    for line in lines:
        text = NSAttributedString.alloc().initWithString_attributes_(line, {CoreText.kCTFontAttributeName: font})
        Quartz.CGContextSetTextPosition(context, 50, y)
        CoreText.CTLineDraw(CoreText.CTLineCreateWithAttributedString(text), context)
        y -= 22
    Quartz.CGPDFContextEndPage(context)
    Quartz.CGPDFContextClose(context)
    return path


def _snapshot(ctx: Ctx, folder: str) -> dict[str, tuple[str, float]]:
    """relative path -> (sha256, mtime) for every file under a sandbox folder."""
    base = ctx.p(folder)
    return {str(p.relative_to(base)): (sb.sha(p), p.stat().st_mtime) for p in sb.files_under(base)}


def _changed(ctx: Ctx, folder: str, before: dict, names=None) -> list[str]:
    """Files of `before` (or only `names`) that are gone, or whose content or mtime changed."""
    base = ctx.p(folder)
    bad = []
    for rel, (digest, mtime) in before.items():
        if names is not None and rel not in names:
            continue
        path = base / rel
        if not path.exists():
            bad.append(f"{rel} is gone")
        elif sb.sha(path) != digest:
            bad.append(f"{rel} was changed")
        elif abs(path.stat().st_mtime - mtime) > 1:
            bad.append(f"{rel} was re-saved (mtime changed)")
    return bad


# --- spreadsheets whose formulas have no cached values (openpyxl, Mint's edit_spreadsheet) -------------

_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")


def _col(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


class Grid:
    """One sheet's cells, with simple formulas (SUM, cell refs, + - * /, ROUND) worked out."""

    def __init__(self, values: dict, formulas: dict, title: str):
        self.values, self.formulas, self.title = values, formulas, title

    def value(self, row: int, col: int, depth: int = 0):
        formula = self.formulas.get((row, col))
        cached = self.values.get((row, col))
        if not (isinstance(formula, str) and formula.startswith("=")) or cached is not None or depth > 20:
            return cached
        return self._eval(formula[1:], depth + 1)

    def _num(self, row: int, col: int, depth: int) -> float:
        got = self.value(row, col, depth)
        return got if isinstance(got, (int, float)) and not isinstance(got, bool) else 0.0

    def _range(self, a: str, b: str, depth: int) -> list[float]:
        (c1, r1), (c2, r2) = [(_col(m.group(1)), int(m.group(2))) for m in (_REF.fullmatch(a), _REF.fullmatch(b))]
        return [self._num(r, c, depth) for r in range(min(r1, r2), max(r1, r2) + 1)
                for c in range(min(c1, c2), max(c1, c2) + 1)]

    def _eval(self, expr: str, depth: int):
        expr = expr.upper().replace("$", "")
        span = r"([A-Z]{1,3}\d+):([A-Z]{1,3}\d+)"
        expr = re.sub(r"SUM\(([^()]*)\)", lambda m: "(" + "+".join(
            str(sum(self._range(*r.groups(), depth))) if (r := re.fullmatch(span, part.strip())) else part
            for part in m.group(1).split(",")) + ")", expr)
        expr = re.sub(r"ROUND\(", "round(", expr)
        expr = _REF.sub(lambda m: repr(self._num(int(m.group(2)), _col(m.group(1)), depth)), expr)
        if not re.fullmatch(r"[\d.+\-*/() ,eround]*", expr):
            return None
        try:
            return eval(expr, {"__builtins__": {}}, {"round": round})  # noqa: S307 - digits and operators only
        except Exception:  # noqa: BLE001 - an unworkable formula is just no value
            return None

    def rows(self) -> list[list]:
        if not self.values and not self.formulas:
            return []
        keys = set(self.values) | set(self.formulas)
        last_row, last_col = max(r for r, _ in keys), max(c for _, c in keys)
        out = []
        for r in range(1, last_row + 1):
            row = [self.value(r, c) for c in range(1, last_col + 1)]
            if any(v is not None and str(v).strip() for v in row):
                out.append(row)
        return out


def grids(path: Path) -> dict[str, Grid]:
    import openpyxl
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    out = {}
    for sheet in formulas.worksheets:
        vals = values[sheet.title]
        cells_v = {(c.row, c.column): c.value for r in vals.iter_rows() for c in r if c.value is not None}
        cells_f = {(c.row, c.column): c.value for r in sheet.iter_rows() for c in r if c.value is not None}
        out[sheet.title] = Grid(cells_v, cells_f, sheet.title)
    return out


def number_formats(path: Path) -> str:
    import openpyxl
    book = openpyxl.load_workbook(path)
    return " ".join(c.number_format for s in book.worksheets for r in s.iter_rows() for c in r)


def _row(rows: list[list], word: str) -> list | None:
    word = word.lower()
    return next((r for r in rows if any(word in str(c).lower() for c in r if isinstance(c, str))), None)


def _has(row: list | None, value: float, tol: float = 0.011) -> bool:
    return row is not None and any((n := number(c)) is not None and abs(n - value) <= tol for c in row)


def _new_xlsx(ctx: Ctx, folder: Path) -> list[Path]:
    return [p for p in sb.new_files(ctx.started, (".xlsx",), (folder,)) if not p.name.startswith("~$")]


# ============================================================================================================
# 1. hard-shop-cheapest: 4 web pages -> spreadsheet with a Total -> 2-line summary
# ============================================================================================================

SHOP_PICK = [("Seed Tray", 3.15), ("Bamboo Spatula", 3.49), ("Sticky Notes", 3.75)]
SHOP_TRAPS = ["Whisk", "Plant Labels", "Paper Clips", "Twine"]    # out of stock, pre-order, 4th cheapest


@setup
def hard_shop(ctx: Ctx):
    ctx.p("Shop").mkdir(parents=True, exist_ok=True)


@check
def hard_shop_ok(ctx: Ctx):
    path = ctx.p("Shop", "cheapest.xlsx")
    if not path.exists():
        found = [p.name for p in sb.new_files(ctx.started, (".xlsx",))]
        return fail("no Shop/cheapest.xlsx" + (f" (new sheets: {found})" if found else ""))
    rows = next(iter(grids(path).values())).rows()
    total = _row(rows, "total")
    items = [r for r in rows if r is not total and any(number(c) is not None for c in r)]
    names = [next((n for n, _ in SHOP_PICK if _row([r], n)), None) for r in items]
    want = [n for n, _ in SHOP_PICK]
    if names != want:
        shown = [" ".join(str(c) for c in r if c is not None) for r in items]
        return fail(f"product rows are {shown}, want {want} (cheapest in-stock first)")
    wrong = [n for (n, price), r in zip(SHOP_PICK, items) if not _has(r, price)]
    if wrong:
        return fail(f"wrong price for {wrong}")
    traps = [t for t in SHOP_TRAPS if _row(rows, t)]
    if traps:
        return fail(f"the sheet lists {traps} (out of stock / pre-order / not in the 3 cheapest)")
    if not _has(total, 10.39):
        return fail(f"no Total of 10.39 (total row: {total})")
    summary = _text(ctx.p("Shop", "summary.md"))
    if summary is None:
        return fail("the sheet is right, but there is no Shop/summary.md")
    lines = _lines(summary)
    if len(lines) != 2:
        return fail(f"summary.md has {len(lines)} lines, want 2")
    if "seed tray" not in lines[0].lower():
        return fail(f"summary line 1 does not name the Seed Tray: {lines[0][:80]!r}")
    if "10.39" not in lines[1]:
        return fail(f"summary line 2 does not give the total 10.39: {lines[1][:80]!r}")
    return ok("3 cheapest in-stock products in order, Total 10.39, 2-line summary")


# ============================================================================================================
# 2. hard-branch-404: the named page is a 404; the branch list has the renamed page
# ============================================================================================================

_TIME = re.compile(r"\b(\d{1,2})(?:[:.](\d{2}))?\s*([ap])\.?\s*m\b\.?|\b(\d{1,2})[:.](\d{2})\b", re.I)


def _times(text: str) -> set[str]:
    found = set()
    for m in _TIME.finditer(text):
        if m.group(3):
            hour, minute = int(m.group(1)) % 12 + (12 if m.group(3).lower() == "p" else 0), int(m.group(2) or 0)
        else:
            hour, minute = int(m.group(4)), int(m.group(5))
        found.add(f"{hour:02d}:{minute:02d}")
    return found


@check
def hard_branch_ok(ctx: Ctx):
    text = _text(ctx.p("hours.txt"))
    if text is None:
        return fail("no hours.txt")
    got = _times(text)
    want = {"08:30", "18:00", "10:00", "16:00"}
    others = {"09:00", "17:30", "13:00", "07:00", "19:00", "11:00", "14:00"} & got
    if others:
        return fail(f"hours.txt has another branch's times {sorted(others)} (Riverdale/Hillcrest?)")
    if not want <= got:
        return fail(f"hours.txt lacks {sorted(want - got)} (has {sorted(got)})")
    if "closed" not in text.lower():
        return fail("hours.txt does not say Sunday is closed")
    return ok("Riverside Quay hours (found via the branch list after the 404)")


# ============================================================================================================
# 3. hard-sales-clean: mixed date formats, currency symbols, duplicates written differently
# ============================================================================================================

SALES_RAW = """order,date,customer,amount
1001,2026-09-03,Northwind,"$1,200.50"
1002,23/09/2026,Contoso ,USD 300
1003,"Sep 5, 2026",Fabrikam,75

1002,2026-09-23,Contoso,300.00
1004,5 Sept 2026, Northwind, $ 20.5
1006,17.09.2026,Fabrikam,$89.99
1005,2026/09/14,Tailspin,"1,050"
1001,2026-09-03,Northwind,"$1,200.50"
"""
SALES_CLEAN = [["1001", "2026-09-03", "Northwind", "1200.50"], ["1002", "2026-09-23", "Contoso", "300.00"],
               ["1003", "2026-09-05", "Fabrikam", "75.00"], ["1004", "2026-09-05", "Northwind", "20.50"],
               ["1005", "2026-09-14", "Tailspin", "1050.00"], ["1006", "2026-09-17", "Fabrikam", "89.99"]]
SALES_TOTAL = "2735.99"


@setup
def hard_sales(ctx: Ctx):
    ctx.data["raw"] = sb.sha(sb.write("Data/sales-raw.csv", SALES_RAW))


@check
def hard_sales_ok(ctx: Ctx):
    raw = ctx.p("Data", "sales-raw.csv")
    if not raw.exists() or sb.sha(raw) != ctx.data["raw"]:
        return fail("sales-raw.csv was changed or removed (the user said not to)")
    path = ctx.p("Data", "sales-clean.csv")
    if not path.exists():
        return fail("no Data/sales-clean.csv")
    rows = _csv(path)
    if not rows or [c.lower() for c in rows[0]] != ["order", "date", "customer", "amount"]:
        return fail(f"header is {rows[0] if rows else None}, want order,date,customer,amount")
    body = rows[1:]
    if body != SALES_CLEAN:
        if len(body) != len(SALES_CLEAN):
            return fail(f"{len(body)} data rows, want {len(SALES_CLEAN)} (one per order)")
        diff = [f"{got} != {want}" for got, want in zip(body, SALES_CLEAN) if got != want]
        return fail("rows differ: " + "; ".join(diff[:3]))
    total = _text(ctx.p("Data", "total.txt"))
    if total is None:
        return fail("clean CSV right, but no Data/total.txt")
    if SALES_TOTAL not in total.replace(",", ""):
        return fail(f"total.txt says {total.strip()[:60]!r}, want {SALES_TOTAL}")
    return ok("6 orders, ISO dates, plain 2-decimal amounts, sorted; total 2735.99; raw file untouched")


# ============================================================================================================
# 4. hard-recon-pdf-csv: a supplier PDF and our ledger CSV disagree
# ============================================================================================================

STATEMENT = [("NW-2201", "2026-09-02", 540.00), ("NW-2202", "2026-09-06", 1275.40), ("NW-2203", "2026-09-11", 89.90),
             ("NW-2204", "2026-09-15", 2310.00), ("NW-2205", "2026-09-22", 460.25), ("NW-2206", "2026-09-27", 118.00)]
LEDGER = {"NW-2201": 540.00, "NW-2202": 1257.40, "NW-2203": 89.90, "NW-2204": 2310.00, "NW-2205": 406.25}
MISMATCH = {"NW-2202": (1275.40, 1257.40, 18.00), "NW-2205": (460.25, 406.25, 54.00), "NW-2206": (118.00, None, None)}


@setup
def hard_recon(ctx: Ctx):
    lines = ["Northwind Supplies - Statement of account, September 2026", "Customer: MintBench Ltd", "",
             "Invoice        Date            Amount"]
    lines += [f"{n}        {d}      ${a:,.2f}" for n, d, a in STATEMENT]
    lines += ["", f"Total due: ${sum(a for *_, a in STATEMENT):,.2f}", "Please pay within 30 days."]
    ctx.data["pdf"] = sb.sha(_pdf(ctx.p("Recon", "supplier-statement.pdf"), lines))
    ctx.data["csv"] = sb.sha(sb.write("Recon/our-ledger.csv", "invoice,amount\n" +
                                      "".join(f"{n},{a:.2f}\n" for n, a in LEDGER.items())))


@check
def hard_recon_ok(ctx: Ctx):
    path = ctx.p("Recon", "mismatches.csv")
    if not path.exists():
        return fail("no Recon/mismatches.csv")
    rows = _csv(path)
    header = [c.lower() for c in rows[0]] if rows else []
    cols = {name: next((i for i, h in enumerate(header) if name in h), None)
            for name in ("invoice", "statement", "ledger", "difference")}
    if None in cols.values():
        return fail(f"header is {rows[0] if rows else None}, want invoice, statement, ledger, difference")
    got = {}
    for row in rows[1:]:
        row += [""] * (len(header) - len(row))
        got[row[cols["invoice"]].upper()] = row
    if sorted(got) != sorted(MISMATCH):
        return fail(f"listed {sorted(got)}, want {sorted(MISMATCH)} (2 amounts differ, NW-2206 is not in the ledger)")
    for inv, (st, led, diff) in MISMATCH.items():
        row = got[inv]
        cell = {k: row[i].replace("$", "").replace(",", "").strip() for k, i in cols.items()}
        if number(cell["statement"]) is None or abs(number(cell["statement"]) - st) > 0.005:
            return fail(f"{inv}: statement {cell['statement']!r}, want {st:.2f}")
        if led is None:
            if number(cell["ledger"]) not in (None, 0.0):
                return fail(f"{inv} is not in the ledger, but the ledger column says {cell['ledger']!r}")
            if number(cell["difference"]) not in (None, st):
                return fail(f"{inv}: difference {cell['difference']!r} (want empty or {st:.2f})")
            continue
        if number(cell["ledger"]) is None or abs(number(cell["ledger"]) - led) > 0.005:
            return fail(f"{inv}: ledger {cell['ledger']!r}, want {led:.2f}")
        if number(cell["difference"]) is None or abs(number(cell["difference"]) - diff) > 0.005:
            return fail(f"{inv}: difference {cell['difference']!r}, want {diff:.2f} (statement minus ledger)")
    if sb.sha(ctx.p("Recon", "our-ledger.csv")) != ctx.data["csv"]:
        return fail("our-ledger.csv was changed")
    return ok("NW-2202 (+18.00), NW-2205 (+54.00) and NW-2206 (missing from the ledger)")


# ============================================================================================================
# 5. hard-acme-merge: near-duplicate folder names, same file names, a different client to leave alone
# ============================================================================================================

ACME = {  # path -> (content, days ago)
    "Acme/contract.pdf": (lambda: pdf_bytes("Acme contract v2 signed"), 3),
    "Acme/logo.png": (lambda: image_bytes("navy"), 30),
    "acme_old/contract.pdf": (lambda: pdf_bytes("Acme contract v1 draft"), 60),
    "acme_old/brief.txt": (lambda: "Acme brief v1 (first draft)\n", 50),
    "acme_old/2025/notes.txt": (lambda: "Acme kickoff notes, 2025\n", 300),
    "Acme Corp (copy)/brief.txt": (lambda: "Acme brief v2 (approved)\n", 10),
    "Acme Corp (copy)/invoice-0917.pdf": (lambda: pdf_bytes("Acme invoice 0917"), 12),
    "Acme-Widgets/contract.pdf": (lambda: pdf_bytes("Acme-Widgets contract, a different client"), 1),
    "Acme-Widgets/price-list.txt": (lambda: "Acme-Widgets price list\n", 5),
}
ACME_TOP = {"contract.pdf": "Acme/contract.pdf", "logo.png": "Acme/logo.png",
            "brief.txt": "Acme Corp (copy)/brief.txt", "invoice-0917.pdf": "Acme Corp (copy)/invoice-0917.pdf",
            "notes.txt": "acme_old/2025/notes.txt"}
ACME_OLDER = ["acme_old/contract.pdf", "acme_old/brief.txt"]


@setup
def hard_acme(ctx: Ctx):
    ctx.data["sha"] = {rel: sb.sha(sb.write(f"Clients/{rel}", make(), mtime=_ago(days)))
                       for rel, (make, days) in ACME.items()}
    ctx.data["widgets"] = _snapshot(ctx, "Clients/Acme-Widgets")


@check
def hard_acme_ok(ctx: Ctx):
    sha = ctx.data["sha"]
    acme, older = ctx.p("Clients", "Acme"), ctx.p("Clients", "Acme", "older")
    problems = []
    for name, src in ACME_TOP.items():
        path = acme / name
        if not path.exists() or sb.sha(path) != sha[src]:
            problems.append(f"Acme/{name} is not the newest version (from {src})")
    extra = [p.name for p in acme.iterdir() if p.is_file() and not p.name.startswith(".") and p.name not in ACME_TOP] \
        if acme.exists() else []
    if extra:
        problems.append(f"extra files in Acme/: {extra}")
    kept = sorted(sb.sha(p) for p in sb.files_under(older))
    if kept != sorted(sha[r] for r in ACME_OLDER):
        problems.append(f"Acme/older/ holds {[p.name for p in sb.files_under(older)]}, want the old contract and brief")
    for gone in ("acme_old", "Acme Corp (copy)"):
        if ctx.p("Clients", gone).exists():
            problems.append(f"'{gone}' is still there")
    bad = _changed(ctx, "Clients/Acme-Widgets", ctx.data["widgets"])
    if bad or len(sb.files_under(ctx.p("Clients", "Acme-Widgets"))) != 2:
        problems.append(f"Acme-Widgets (another client) was touched: {bad or 'files added'}")
    return fail("; ".join(problems)) if problems else ok("merged, newest kept, older in Acme/older, Widgets untouched")


# ============================================================================================================
# 6. hard-week-digest: only the notes modified in the last 7 days, oldest first; originals untouched
# ============================================================================================================

LOGS = {"standup-a.md": ("ORCHID: shipped the export button", 2), "standup-b.md": ("HERON: fixed the flaky test", 9),
        "standup-c.md": ("LUPIN: started the billing rewrite", 5), "standup-d.md": ("MAPLE: demoed search to sales", 1),
        "standup-e.md": ("OTTER: planned the offsite", 20), "retro.md": ("BIRCH: retro on the outage", 40),
        "plan.txt": ("CEDAR: plan for next quarter", 3)}
LOGS_IN = ["standup-c.md", "standup-a.md", "standup-d.md"]          # oldest first


@setup
def hard_week(ctx: Ctx):
    for name, (line, days) in LOGS.items():
        sb.write(f"Logs/{name}", f"# Team notes\n\n- {line}\n- no blockers\n", mtime=_ago(days))
    ctx.data["logs"] = _snapshot(ctx, "Logs")


@check
def hard_week_ok(ctx: Ctx):
    bad = _changed(ctx, "Logs", ctx.data["logs"])
    if bad:
        return fail("originals touched: " + "; ".join(bad))
    text = _text(ctx.p("Logs", "this-week.md"))
    if text is None:
        return fail("no Logs/this-week.md")
    words = {name: LOGS[name][0].split(":")[0] for name in LOGS}
    wrong = [f"{n} ({w})" for n, w in words.items() if n not in LOGS_IN and w in text]
    if wrong:
        return fail(f"includes notes not modified in the last 7 days / not Markdown: {wrong}")
    missing = [n for n in LOGS_IN if words[n] not in text]
    if missing:
        return fail(f"missing {missing}")
    at = [text.index(words[n]) for n in LOGS_IN]
    if at != sorted(at):
        return fail("not oldest first (want standup-c, standup-a, standup-d)")
    unnamed = [n for n in LOGS_IN if n.rsplit(".", 1)[0] not in text]
    if unnamed:
        return fail(f"no heading with the file name for {unnamed}")
    return ok("3 recent notes, oldest first, under their names; originals untouched")


# ============================================================================================================
# 7. hard-spanish-brief: in Spanish, under 100 words, version and dates exact
# ============================================================================================================

RELEASE = """# Lumen 4.12.0 release notes

Released on 2026-09-21.

## New
- **Offline mode for the mobile app.** Lists, notes and reports now open without a connection and sync
  when you are back online. Changes made offline are merged, never overwritten.
- **CSV export for every report.** Any report can now be downloaded as a CSV file from the Share menu, with
  the same filters and column order you see on screen.
- **Faster search.** Search results now appear about 30% faster on large workspaces, thanks to a new index
  that is built in the background.

## Fixed
- A login loop on Safari when third-party cookies were blocked.
- Reminders that showed the wrong time after a daylight-saving change.
- The dashboard chart legend overlapping the numbers on small screens.

## Changing soon
The old v1 API will be switched off on 2026-12-01. Please move your integrations to the v2 API before then;
the migration guide explains every endpoint and gives examples for the three most common integrations.

## Thanks
Thank you to everyone who reported bugs in the beta, especially the teams who tested offline mode on long
flights and train journeys. Your feedback shaped this release.
"""
_ES = {"el", "la", "los", "las", "de", "del", "y", "en", "para", "con", "una", "un", "que", "se", "por", "más",
       "ahora", "nueva", "nuevo", "sin", "antes", "al", "es"}
_EN = {"the", "and", "with", "for", "will", "is", "of", "to"}


@setup
def hard_spanish(ctx: Ctx):
    ctx.data["src"] = sb.sha(sb.write("Docs/release-notes.md", RELEASE))


@check
def hard_spanish_ok(ctx: Ctx):
    text = _text(ctx.p("Docs", "resumen.md"))
    if text is None:
        return fail("no Docs/resumen.md")
    words = [w for w in re.findall(r"[\wÀ-ÿ][\wÀ-ÿ.'’/-]*", text) if re.search(r"\w", w)]
    if len(words) >= 100:
        return fail(f"{len(words)} words (want under 100)")
    if len(words) < 25:
        return fail(f"only {len(words)} words: not a summary")
    lower = [w.lower().strip(".") for w in words]
    es, en = {w for w in lower if w in _ES}, [w for w in lower if w in _EN]
    if len(es) < 6 or len(en) > 3:
        return fail(f"does not read as Spanish (Spanish words {sorted(es)}, English {en[:8]})")
    missing = [s for s in ("4.12.0", "2026-09-21", "2026-12-01") if s not in text]
    if missing:
        return fail(f"lacks {missing} written exactly")
    if "csv" not in text.lower() or "api" not in text.lower():
        return fail("leaves out the CSV export or the API shutdown")
    if sb.sha(ctx.p("Docs", "release-notes.md")) != ctx.data["src"]:
        return fail("release-notes.md was changed")
    return ok(f"Spanish, {len(words)} words, version and dates exact")


# ============================================================================================================
# 8. hard-trip-euros: 3 turns (make, add a row, "actually use euros")
# ============================================================================================================

TRIP_USD = {"flight": 420, "hotel": 285, "food": 160, "museum": 29, "taxi": 35}
RATE = 0.92


def _trip(ctx: Ctx, rate: float, want_eur: bool):
    folder = ctx.p("Trip")
    candidates = [ctx.p("Trip", "lisbon.xlsx")] + _new_xlsx(ctx, folder)
    reasons = []
    for path in dict.fromkeys(candidates):
        if not path.exists():
            continue
        rows = next(iter(grids(path).values())).rows()
        missing = [f"{k} {round(v * rate, 2)}" for k, v in TRIP_USD.items() if not _has(_row(rows, k), round(v * rate, 2))]
        if missing:
            reasons.append(f"{path.name}: no {missing}")
            continue
        total = _row(rows, "total")
        want = round(sum(TRIP_USD.values()) * rate, 2)
        if not _has(total, want):
            reasons.append(f"{path.name}: no Total of {want}")
            continue
        if rows.index(_row(rows, "taxi")) > rows.index(total):
            reasons.append(f"{path.name}: the taxi row is below the Total")
            continue
        if want_eur:
            shown = " ".join(str(c) for r in rows for c in r if isinstance(c, str)) + " " + number_formats(path)
            if not re.search(r"EUR|€|euro", shown, re.I):
                reasons.append(f"{path.name}: nothing says EUR")
                continue
        return ok(f"{path.name}: 5 items + Total {want}" + (" in EUR" if want_eur else ""))
    return fail("; ".join(reasons) or "no Trip/lisbon.xlsx")


@setup
def hard_trip(ctx: Ctx):
    ctx.p("Trip").mkdir(parents=True, exist_ok=True)


@check
def hard_trip_usd(ctx: Ctx):
    return _trip(ctx, 1.0, False)


@check
def hard_trip_ok(ctx: Ctx):
    return _trip(ctx, RATE, True)


# ============================================================================================================
# 9. hard-second-draft: 4 turns ("the second one", "that file", "shorter", a title)
# ============================================================================================================

_FILLER = ("It is built for small teams who are tired of juggling tools. Setup takes minutes and needs no "
           "training. Customers tell us they save hours every week. ")


def _pitch(product: str, price: str, size: int) -> str:
    head = f"{product} is the simplest way to run your team's week, for just {price} a month. "
    text = head
    while len(text) < size:
        text += _FILLER
    return text[:size].rsplit(". ", 1)[0] + ".\n"


PITCHES = {"draft-a.md": ("Kelvo", "$19", 900), "draft-b.md": ("Tamsin", "$99", 2400),
           "draft-c.md": ("Orbitra", "$49", 1600), "draft-d.md": ("Pruna", "$9", 300)}


@setup
def hard_second(ctx: Ctx):
    for name, (product, price, size) in PITCHES.items():
        sb.write(f"Pitches/{name}", _pitch(product, price, size))
    ctx.data["pitches"] = _snapshot(ctx, "Pitches")


@check
def hard_second_ok(ctx: Ctx):
    bad = _changed(ctx, "Pitches", ctx.data["pitches"])
    if bad:
        return fail("originals touched: " + "; ".join(bad))
    text = _text(ctx.p("Pitches", "pitch-short.md"))
    if text is None:
        return fail("no Pitches/pitch-short.md")
    lines = _lines(text)
    if not lines or lines[0] != "# Pitch":
        return fail(f"first line is {lines[0][:40] if lines else None!r}, want '# Pitch'")
    body = " ".join(lines[1:])
    others = [p for p, *_ in PITCHES.values() if p != "Orbitra" and p in body]
    if "Orbitra" not in body:
        return fail(f"not the second-biggest draft (Orbitra, draft-c.md); names {others}")
    if others:
        return fail(f"mentions other products {others}")
    if "49" not in body:
        return fail("the price ($49) was dropped")
    sentences = [s for s in re.split(r"(?<=[.!?])\s+", body) if re.search(r"\w", s)]
    if len(sentences) > 3:
        return fail(f"{len(sentences)} sentences, want at most 3")
    extra = sorted({p.name for p in sb.files_under(ctx.p("Pitches"))} - set(PITCHES) - {"pitch-short.md"})
    if extra:
        return fail(f"extra files in Pitches: {extra}")
    return ok(f"'# Pitch' + {len(sentences)} sentences about Orbitra at $49; drafts untouched")


# ============================================================================================================
# 10. hard-invoice-dated: turn 1 uses the save date; turn 2 corrects it to the printed invoice date
# ============================================================================================================

INVOICES2 = {  # name -> (printed date, saved)
    "inv-01.txt": ("3 September 2026", dt.datetime(2026, 9, 4, 10)),
    "inv-02.txt": ("28 Aug 2026", dt.datetime(2026, 9, 10, 10)),
    "inv-03.txt": ("2026-09-12", dt.datetime(2026, 9, 15, 10)),
    "inv-04.txt": ("July 30, 2026", dt.datetime(2026, 9, 20, 10)),
    "inv-05.txt": ("1 Sep 2026", dt.datetime(2026, 8, 25, 10)),
    "inv-06.txt": ("10 August 2026", dt.datetime(2026, 8, 12, 10)),
}
SEPT = ["inv-01.txt", "inv-03.txt", "inv-05.txt"]


@setup
def hard_invoice_dated(ctx: Ctx):
    ctx.data["sha"] = {}
    for n, (name, (printed, saved)) in enumerate(INVOICES2.items(), 1):
        body = (f"INVOICE {name[:-4].upper()}\nFrom: Lakeside Print Co.\nInvoice date: {printed}\n\n"
                f"Printing ........ {n * 37.5:.2f}\nTotal due: ${n * 37.5:.2f}\n")
        ctx.data["sha"][name] = sb.sha(sb.write(f"Invoices/{name}", body, mtime=saved))


@check
def hard_invoice_dated_ok(ctx: Ctx):
    sha = ctx.data["sha"]
    sept = {p.name: sb.sha(p) for p in sb.files_under(ctx.p("Invoices", "2026-09"))}
    folder = ctx.p("Invoices")
    top = {p.name: sb.sha(p) for p in (folder.iterdir() if folder.exists() else [])
           if p.is_file() and not p.name.startswith(".")}
    if sorted(sept) != SEPT:
        return fail(f"2026-09/ holds {sorted(sept)}, want {SEPT} (invoices DATED September)")
    rest = sorted(set(INVOICES2) - set(SEPT))
    if sorted(top) != rest:
        return fail(f"Invoices/ top level holds {sorted(top)}, want {rest}")
    changed = [n for n, h in {**sept, **top}.items() if sha[n] != h]
    return fail(f"invoice content changed: {changed}") if changed else ok("exactly the 3 invoices dated September")


# ============================================================================================================
# 11. hard-undo-replace: replace all, undo, then replace only the heading
# ============================================================================================================

QPLAN = """# Q3 plan

Goals for Q3:
- Ship the Q3 pricing page
- Hire two engineers before the end of Q3

Owner: Dana (Q3 lead)
"""


@setup
def hard_undo_replace(ctx: Ctx):
    ctx.data["orig"] = sb.sha(sb.write("Plan/q-plan.md", QPLAN))


@check
def hard_undo_restored(ctx: Ctx):
    text = _text(ctx.p("Plan", "q-plan.md"))
    if text is None:
        return fail("q-plan.md is gone after the undo")
    return ok("undo put the file back") if _lines(text) == _lines(QPLAN) else fail(
        f"after 'undo that' the file is not the original ({text.count('Q4')} x Q4)")


@check
def hard_undo_replace_ok(ctx: Ctx):
    text = _text(ctx.p("Plan", "q-plan.md"))
    if text is None:
        return fail("q-plan.md is gone")
    want = _lines(QPLAN.replace("# Q3 plan", "# Q4 plan", 1))
    got = _lines(text)
    if got == want:
        return ok("only the heading says Q4; the other 4 Q3s kept")
    return fail(f"file has {text.count('Q4')} x Q4 and {text.count('Q3')} x Q3 (want 1 and 4); first line {got[:1]}")


# ============================================================================================================
# 12. hard-undo-move-copy: move, "undo that", copy under a new name instead
# ============================================================================================================

@setup
def hard_undo_move(ctx: Ctx):
    ctx.data["report"] = sb.sha(sb.write("Share/report-final.pdf", pdf_bytes("MintBench final report 2026-09")))
    sb.write("Share/readme.txt", "Shared with the board.\n")
    ctx.data["old"] = sb.sha(sb.write("Archive/2025-report.pdf", pdf_bytes("2025 report")))


@check
def hard_undo_move_ok(ctx: Ctx):
    report = ctx.data["report"]
    share = ctx.p("Share", "report-final.pdf")
    if not share.exists() or sb.sha(share) != report:
        return fail("Share/report-final.pdf is not back (undo)")
    copy = ctx.p("Archive", "report-final-2026-09.pdf")
    if not copy.exists() or sb.sha(copy) != report:
        return fail("no Archive/report-final-2026-09.pdf copy")
    archive = sorted(p.name for p in sb.files_under(ctx.p("Archive")))
    if archive != ["2025-report.pdf", "report-final-2026-09.pdf"]:
        return fail(f"Archive holds {archive} (the moved original should be gone)")
    if sb.sha(ctx.p("Archive", "2025-report.pdf")) != ctx.data["old"]:
        return fail("2025-report.pdf changed")
    if not ctx.p("Share", "readme.txt").exists():
        return fail("Share/readme.txt is gone")
    return ok("original back in Share, one renamed copy in Archive")


# ============================================================================================================
# 13. hard-specs-locked: two of five files can't be changed partway through
# ============================================================================================================

SPECS = {f"spec-{c}.md": f"# Spec {c.upper()}\n\nScope: part {c.upper()} of the MintBench app.\n" for c in "abcde"}
LOCKED = ["spec-c.md", "spec-e.md"]
STAMP = "Reviewed: 2026-09-29"


@setup
def hard_specs(ctx: Ctx):
    for name, text in SPECS.items():
        sb.write(f"Specs/{name}", text)
    subprocess.run(["chflags", "uchg", str(ctx.p("Specs", "spec-c.md"))], check=True)     # Finder's "Locked"
    os.chmod(ctx.p("Specs", "spec-e.md"), 0o444)                                         # read-only
    ctx.data["specs"] = _snapshot(ctx, "Specs")


@check
def hard_specs_ok(ctx: Ctx):
    problems = []
    for name, text in SPECS.items():
        path = ctx.p("Specs", name)
        if name in LOCKED:
            continue
        got = _lines(_text(path) or "")
        if got != _lines(text) + [STAMP]:
            problems.append(f"{name} ends with {got[-1:] or 'nothing'} (want its text + '{STAMP}')")
    bad = _changed(ctx, "Specs", ctx.data["specs"], LOCKED)
    if bad:
        problems.append("forced a protected file: " + ", ".join(bad))
    c, e = ctx.p("Specs", "spec-c.md"), ctx.p("Specs", "spec-e.md")
    if c.exists() and not c.stat().st_flags & stat.UF_IMMUTABLE:
        problems.append("spec-c.md was unlocked")
    if e.exists() and e.stat().st_mode & stat.S_IWUSR:
        problems.append("spec-e.md was made writable")
    skipped = _text(ctx.p("Specs", "skipped.txt"))
    if skipped is None:
        problems.append("no Specs/skipped.txt")
    else:
        listed = sorted(set(re.findall(r"spec-[a-e]", skipped)))
        if listed != ["spec-c", "spec-e"]:
            problems.append(f"skipped.txt lists {listed}, want spec-c and spec-e")
    return fail("; ".join(problems)) if problems else ok("3 stamped, 2 protected files skipped and listed")


# ============================================================================================================
# 14. hard-astra-grinders: Astra researches and saves; Mint then edits that file
# ============================================================================================================

GRINDERS = {"Burrline": (189, 40), "Ember": (129, 18), "Tidewell": (159, 55)}


@check
def hard_astra_ok(ctx: Ctx):
    path = ctx.p("Research", "grinders.md")
    if not path.exists():
        found = [str(p) for p in sb.new_files(ctx.started, (".md",)) if "grinder" in p.name.lower()]
        return fail("no Research/grinders.md" + (f" (found {found})" if found else ""))
    text = path.read_text(errors="replace")
    table = [line for line in text.splitlines() if line.count("|") >= 3]
    for name, (price, settings) in GRINDERS.items():
        row = next((line for line in table if name.lower() in line.lower()), None)
        if row is None:
            return fail(f"no table row for {name}")
        nums = _numbers(row)
        if price not in nums or settings not in nums:
            return fail(f"{name} row lacks price {price} or {settings} settings: {row.strip()[:100]!r}")
        if "steel" not in row.lower() and "ceramic" not in row.lower():
            return fail(f"{name} row lacks the burr material")
    last = _lines(text)[-1]
    if not last.lower().lstrip("*_-# ").startswith("pick"):
        return fail(f"the last line is {last[:80]!r}, want 'Pick: ...' (added after Astra saved the file)")
    picked = [n for n in GRINDERS if n.lower() in last.lower()]
    if picked != ["Tidewell"]:
        return fail(f"Pick line names {picked}, want Tidewell (cheapest with steel burrs)")
    return ok("Astra's table with 3 grinders + 'Pick: Tidewell Pro' at the end")


# ============================================================================================================
# 15. hard-clip-three: paste the last three copies, in order, below what is there
# ============================================================================================================

CLIPS = ["MintBench old note: ignore this one", "MintBench step 1: back up the database",
         "MintBench step 2: run the migration", "MintBench step 3: restart the workers"]


@setup
def hard_clip(ctx: Ctx):
    import time
    sb.write("notes.md", "# Deploy notes\n")
    for text in CLIPS:
        sb.clipboard_set(text)
        time.sleep(2.5)                # Mint's clipboard watcher must see each copy


@check
def hard_clip_ok(ctx: Ctx):
    text = _text(ctx.p("notes.md"))
    if text is None:
        return fail("notes.md is gone")
    got = [_bullet(line) for line in _lines(text)]
    want = ["# Deploy notes", *CLIPS[1:]]
    if got == want:
        return ok("heading kept, then steps 1, 2, 3")
    if "# Deploy notes" not in got:
        return fail("the existing '# Deploy notes' line was lost")
    if CLIPS[0] in got:
        return fail("pasted the older 4th copy too")
    return fail(f"notes.md lines are {got}")


# ============================================================================================================
# 16. hard-inventory-edit: set a cell, add a row, a Total for one column, another sheet left alone
# ============================================================================================================

STOCK = [("Item", "Stock", "Price"), ("Blue Mug", 12, 8.5), ("Red Mug", 7, 8.5), ("Teapot", 3, 24.0),
         ("Coaster Set", 25, 6.0)]
NOTES_SHEET = [("Supplier", "Kiln & Co"), ("Reorder when stock is below", 5), ("Last count", "2026-09-20")]
STOCK_WANT = {"Blue Mug": (14, 8.5), "Red Mug": (7, 8.5), "Teapot": (3, 24.0), "Coaster Set": (25, 6.0),
              "Green Mug": (20, 8.5)}


@setup
def hard_inventory(ctx: Ctx):
    import openpyxl
    book = openpyxl.Workbook()
    book.active.title = "Stock"
    for row in STOCK:
        book.active.append(row)
    notes = book.create_sheet("Notes")
    for row in NOTES_SHEET:
        notes.append(row)
    ctx.p("Shop").mkdir(parents=True, exist_ok=True)
    book.save(ctx.p("Shop", "inventory.xlsx"))


@check
def hard_inventory_ok(ctx: Ctx):
    path = ctx.p("Shop", "inventory.xlsx")
    if not path.exists():
        return fail("inventory.xlsx is gone")
    sheets = grids(path)
    if "Stock" not in sheets or "Notes" not in sheets:
        return fail(f"sheets are {list(sheets)}, want Stock and Notes")
    notes = [tuple(r) for r in sheets["Notes"].rows()]
    if notes != NOTES_SHEET:
        return fail(f"the Notes sheet changed: {notes}")
    rows = sheets["Stock"].rows()
    if not rows or [str(c).strip() for c in rows[0][:3]] != ["Item", "Stock", "Price"]:
        return fail(f"the header row changed: {rows[:1]}")
    items = {str(r[0]).strip(): r for r in rows[1:]
             if r and isinstance(r[0], str) and not r[0].strip().lower().startswith("total")}
    problems = []
    for name, (stock, price) in STOCK_WANT.items():
        row = items.get(name)
        if row is None:
            problems.append(f"no {name} row")
        elif number(row[1]) != stock or number(row[2]) is None or abs(number(row[2]) - price) > 0.001:
            problems.append(f"{name} is {row[1:3]}, want [{stock}, {price}]")
    total = next((r for r in rows if r and str(r[0]).strip().lower().startswith("total")), None)
    if total is None:
        problems.append("no Total row")
    else:
        if rows.index(total) != len(rows) - 1:
            problems.append("the Total row is not at the bottom")
        if number(total[1]) != 69:
            problems.append(f"Total stock is {total[1]!r}, want 69")
    extra = sorted(set(items) - set(STOCK_WANT))
    if extra:
        problems.append(f"unexpected rows {extra}")
    return fail("; ".join(problems)) if problems else ok("Blue Mug 14, Green Mug added, Total stock 69, Notes kept")


# ============================================================================================================
# 17. hard-video-chain: cut 0:05-0:15 -> audio of the clip -> durations file; the source untouched
# ============================================================================================================

TONES = "if(lt(t,5),300,if(lt(t,15),700,1100))"      # the wanted part is the only 700 Hz stretch


@setup
def hard_video(ctx: Ctx):
    path = ctx.p("Video", "talk.mp4")
    path.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", "testsrc2=size=1280x720:rate=30",
                    "-f", "lavfi", "-i", f"aevalsrc='0.5*sin(2*PI*{TONES}*t)':s=44100", "-t", "30",
                    "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", str(path)],
                   check=True, capture_output=True)
    ctx.data["talk"] = sb.sha(path)


def tone(path: Path, start: float, seconds: float) -> float:
    """The dominant frequency (zero crossings) of the audio between start and start + seconds."""
    raw = subprocess.run(["ffmpeg", "-v", "error", "-ss", str(start), "-t", str(seconds), "-i", str(path),
                          "-vn", "-ac", "1", "-ar", "8000", "-f", "s16le", "-"], capture_output=True).stdout
    samples = memoryview(raw).cast("h") if len(raw) % 2 == 0 else memoryview(raw[:-1]).cast("h")
    if len(samples) < 800:
        return 0.0
    crossings = sum(1 for a, b in zip(samples, samples[1:]) if (a < 0) != (b < 0))
    return crossings / 2 / (len(samples) / 8000)


def _seconds(text: str) -> list[float]:
    found = []
    for m in re.finditer(r"\b(?:(\d+):)?(\d{1,2}):(\d{2}(?:\.\d+)?)\b", text):
        found.append(int(m.group(1) or 0) * 3600 + int(m.group(2)) * 60 + float(m.group(3)))
    for m in re.finditer(r"(?<![\d:.])(\d+(?:\.\d+)?)(?![\d:])", re.sub(r"\b\d+:\d{2}(?::\d{2})?(?:\.\d+)?", " ", text)):
        found.append(float(m.group(1)))
    return found


@check
def hard_video_ok(ctx: Ctx):
    talk = ctx.p("Video", "talk.mp4")
    if not talk.exists() or sb.sha(talk) != ctx.data["talk"]:
        return fail("talk.mp4 was changed or removed")
    clip, audio = ctx.p("Video", "clip.mp4"), ctx.p("Video", "clip.mp3")
    if not clip.exists():
        return fail(f"no Video/clip.mp4 (Video holds {[p.name for p in sb.files_under(ctx.p('Video'))]})")
    info = sb.media(clip)
    if not info.get("has_video") or abs(info.get("duration", 0) - 10) > 1:
        return fail(f"clip.mp4 is {info.get('duration', 0):.1f} s, want 10 s of video")
    hz = tone(clip, 2, 6)
    if abs(hz - 700) > 60:
        return fail(f"clip.mp4 is not 0:05-0:15 of talk.mp4 (its audio is {hz:.0f} Hz, that part is 700 Hz)")
    if not audio.exists():
        return fail("no Video/clip.mp3")
    info = sb.media(audio)
    if not info.get("has_audio") or abs(info.get("duration", 0) - 10) > 1:
        return fail(f"clip.mp3 is not 10 s of audio ({info})")
    if abs(tone(audio, 2, 6) - 700) > 60:
        return fail("clip.mp3 is not the clip's audio")
    text = _text(ctx.p("Video", "durations.txt"))
    if text is None:
        return fail("no Video/durations.txt")
    tens = [s for s in _seconds(text) if 9 <= s <= 11]
    if len(tens) < 2 or "mp3" not in text.lower() or "mp4" not in text.lower():
        return fail(f"durations.txt does not give both durations (~10 s): {text.strip()[:120]!r}")
    return ok("10 s clip of the right part, its mp3, durations written; talk.mp4 untouched")


# ============================================================================================================
# 18. hard-accountant-memory: remember a person, then use the memory with files in the next turn
# ============================================================================================================

RECEIPT_TOTALS = {"r-0902.txt": 12.40, "r-0909.txt": 58.00, "r-0915.txt": 7.95, "r-0921.txt": 131.10}


@setup
def hard_accountant(ctx: Ctx):
    for name, total in RECEIPT_TOTALS.items():
        sb.write(f"Receipts/{name}", f"RECEIPT {name[2:6]}\nThank you for shopping\n\nTOTAL  ${total:.2f}\n")
    sb.write("Receipts/old-summary.csv", "month,total\nAugust,180.00\n")


@check
def hard_accountant_ok(ctx: Ctx):
    if not sb.bank_matching(("farida",)):
        return fail("Farida (the accountant) was not saved to memory")
    text = _text(ctx.p("Receipts", "for-accountant.md"))
    if text is None:
        return fail("no Receipts/for-accountant.md")
    if "farida" not in text.lower():
        return fail("the note does not greet Farida (from memory)")
    flat = text.replace(",", "")
    if "209.45" not in flat:
        return fail(f"the note lacks the total 209.45 (numbers: {_numbers(text)[:8]})")
    if "180" in flat:
        return fail("used old-summary.csv's 180.00")
    if not re.search(r"\b(4|four)\b", text, re.I):
        return fail("the note does not say there are 4 receipts")
    return ok("remembered Farida; note greets her with 4 receipts totalling 209.45")


# ============================================================================================================
# 19. hard-reminders-csv: a messy task list -> reminders for only my open tasks, due at 9 am
# ============================================================================================================

def _task_rows() -> list[tuple[str, str, str, int]]:
    return [("MintBench: renew the domain", "Sam", "open", 1), ("MintBench: send the Q3 deck", " sam", "Open ", 2),
            ("MintBench: book the venue", "Priya", "open", 1), ("MintBench: pay the invoice", "Sam", "done", 1),
            ("MintBench: order toner", "SAM", "OPEN", 3), ("MintBench: renew the domain", "Sam", "open", 1)]


REMINDERS_WANT = {"domain": 1, "deck": 2, "toner": 3}


@setup
def hard_reminders(ctx: Ctx):
    ap.reminders_ensure(ctx.apple)
    today = dt.date.today()
    lines = ["task,owner,status,due"] + [f"{t},{o},{s},{today + dt.timedelta(days=d):%Y-%m-%d}"
                                         for t, o, s, d in _task_rows()]
    sb.write("tasks.csv", "\n".join(lines) + "\n")


@check
def hard_reminders_ok(ctx: Ctx):
    rows = ap.reminders_list(ctx.apple)
    names = [r["name"] for r in rows]
    for word in ("venue", "invoice"):
        if any(word in n.lower() for n in names):
            return fail(f"made a reminder for the '{word}' task (not Sam's, or done)")
    for word, days in REMINDERS_WANT.items():
        found = [r for r in rows if word in r["name"].lower()]
        if len(found) != 1:
            return fail(f"{len(found)} reminders for '{word}' (want 1; Mint Bench has {names})")
        due = ap.parse_iso(found[0]["due"])
        want = dt.datetime.combine(dt.date.today() + dt.timedelta(days=days), dt.time(9))
        if due is None or due.replace(second=0) != want:
            return fail(f"'{word}' is due {found[0]['due'] or 'never'}, want {want:%Y-%m-%d} 09:00")
    if len(rows) != 3:
        return fail(f"{len(rows)} reminders in Mint Bench, want 3: {names}")
    return ok("3 reminders (Sam's open tasks), each due at 9:00 on its date")


# ============================================================================================================
# 20. hard-rates-convert: a rate from a web page (not last month's) applied to a CSV; original kept
# ============================================================================================================

PRICES_USD = [("A-100", "Desk Lamp", 49.99), ("A-101", "Office Chair", 189.00), ("A-102", "Monitor Arm", 74.00),
              ("A-103", "Cable Tray", 18.25)]
GBP = 0.79


@setup
def hard_rates(ctx: Ctx):
    text = "sku,product,price_usd\n" + "".join(f"{s},{p},{v:.2f}\n" for s, p, v in PRICES_USD)
    ctx.data["usd"] = sb.sha(sb.write("Shop/prices-usd.csv", text))


@check
def hard_rates_ok(ctx: Ctx):
    usd = ctx.p("Shop", "prices-usd.csv")
    if not usd.exists() or sb.sha(usd) != ctx.data["usd"]:
        return fail("prices-usd.csv was changed")
    path = ctx.p("Shop", "prices-gbp.csv")
    if not path.exists():
        return fail("no Shop/prices-gbp.csv")
    rows = _csv(path)
    if not rows or [c.lower() for c in rows[0]] != ["sku", "product", "price_gbp"]:
        return fail(f"header is {rows[0] if rows else None}, want sku,product,price_gbp")
    body = rows[1:]
    if [r[:2] for r in body] != [[s, p] for s, p, _ in PRICES_USD]:
        return fail(f"rows are {[r[:2] for r in body]} (same products, same order wanted)")
    for (sku, _, price), row in zip(PRICES_USD, body):
        want = round(price * GBP + 1e-9, 2)
        cell = row[2] if len(row) > 2 else ""
        if not re.fullmatch(r"\d+\.\d{2}", cell):
            return fail(f"{sku}: {cell!r} is not a plain number with 2 decimals")
        if abs(float(cell) - want) > 0.006:
            old = " (last month's rate 0.81?)" if abs(float(cell) - price * 0.81) < 0.02 else ""
            return fail(f"{sku}: {cell}, want {want:.2f}{old}")
    return ok("4 prices at today's 0.79, 2 decimals; the USD file untouched")
