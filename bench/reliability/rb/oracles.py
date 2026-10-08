"""Oracles: the correct end state for each task, made without Mint (`run.py --selftest`).

A check that fails on its oracle is a broken check. Oracles only touch the sandbox and the
Mint Bench containers, like the tasks themselves.
"""

from __future__ import annotations

import shutil
import subprocess
import urllib.request

from . import apple as ap
from . import sandbox as sb
from .registry import Ctx, next_weekday, tomorrow_at
from .t_files import INBOX, RECEIPTS

ORACLES: dict = {}


def oracle(task_id: str):
    def register(fn):
        ORACLES[task_id] = fn
        return fn
    return register


def _xlsx(path, rows) -> None:
    import openpyxl
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(list(row))
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def _move(src, dst) -> None:
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(src), str(dst))


# --- files ----------------------------------------------------------------------------------------

@oracle("files-organize")
def _(ctx: Ctx):
    for name, (folder, _) in INBOX.items():
        _move(ctx.p("Inbox", name), ctx.p("Inbox", folder, name))


@oracle("files-typo-folder")
def _(ctx: Ctx):
    for name in RECEIPTS:
        if name.startswith("2026-03"):
            _move(ctx.p("Receipts", name), ctx.p("Receipts", "March", name))


@oracle("files-missing-file")
def _(ctx: Ctx):
    with open(ctx.p("notes", "agenda-draft.md"), "a") as f:
        f.write("- Budget review\n")


@oracle("files-locked-file")
def _(ctx: Ctx):
    sb.write("config/settings-light.txt", "theme=light\nfont=14\nsync=on\n")


@oracle("files-dedupe")
def _(ctx: Ctx):
    for name in ("report (1).pdf", "report copy.pdf"):
        ctx.p("Downloads", name).unlink()


@oracle("files-tidy-undo")
def _(ctx: Ctx):
    pass            # the end state (everything back) is the start state; the mid-check is skipped


# --- documents ---------------------------------------------------------------------------------------

@oracle("docs-action-items")
def _(ctx: Ctx):
    sb.write("Docs/action-items.md", "# Action items\n\n- Nina: send the revised budget to finance by Friday\n"
             "- Tom: book the venue for the offsite\n- Lena: draft the hiring plan\n"
             "- Omar: fix the login bug before the release\n")


@oracle("docs-convert-docx")
def _(ctx: Ctx):
    subprocess.run(["textutil", "-convert", "docx", "-output", str(ctx.p("Docs", "proposal.docx")),
                    "-stdin", "-format", "txt"], input=ctx.p("Docs", "proposal.md").read_text(), text=True, check=True)


@oracle("docs-translate")
def _(ctx: Ctx):
    text = ("¡Bienvenido al equipo! Tu primer día empieza a las nueve de la mañana. Por favor, trae tu "
            "tarjeta de identificación a recepción. El almuerzo se ofrece todos los viernes.")
    subprocess.run(["textutil", "-convert", "docx", "-output", str(ctx.p("Docs", "welcome (Spanish).docx")),
                    "-stdin", "-format", "txt"], input=text, text=True, check=True)


def _pdf(path, lines: list[str]) -> None:
    """A one-page PDF with a real text layer (CoreGraphics + CoreText)."""
    import CoreText
    import Quartz
    from Foundation import NSURL, NSAttributedString
    context = Quartz.CGPDFContextCreateWithURL(NSURL.fileURLWithPath_(str(path)), ((0, 0), (612, 792)), None)
    Quartz.CGPDFContextBeginPage(context, None)
    font = CoreText.CTFontCreateWithName("Helvetica", 14, None)
    y = 740
    for line in lines:
        text = NSAttributedString.alloc().initWithString_attributes_(line, {CoreText.kCTFontAttributeName: font})
        Quartz.CGContextSetTextPosition(context, 50, y)
        CoreText.CTLineDraw(CoreText.CTLineCreateWithAttributedString(text), context)
        y -= 24
    Quartz.CGPDFContextEndPage(context)
    Quartz.CGPDFContextClose(context)


@oracle("docs-pdf-report")
def _(ctx: Ctx):
    _pdf(ctx.p("Reports", "MintBench Weekly Report.pdf"),
         ["MintBench Weekly Report", "Wins", "Shipped the new onboarding flow; signed Contoso.", "Risks",
          "The payments vendor may raise prices.", "Next steps", "Start the mobile beta; hire a designer."])


@oracle("docs-textedit-edit")
def _(ctx: Ctx):
    path = ctx.p("Docs", "letter.rtf")
    path.write_text(path.read_text().replace("Dear Sir", "Dear Ms. Rao"))


# --- spreadsheets ---------------------------------------------------------------------------------------

@oracle("sheets-receipts")
def _(ctx: Ctx):
    _xlsx(ctx.p("expenses.xlsx"), [("Vendor", "Date", "Total"), ("Blue Bottle Coffee", "2026-09-02", 4.5),
                                   ("City Cabs", "2026-09-05", 23.8), ("Office Depot", "2026-09-10", 112.35),
                                   ("Pizza Palace", "2026-09-18", 36)])


@oracle("sheets-web-table")
def _(ctx: Ctx):
    _xlsx(sb.storage_folder() / "MintBench fruit prices.xlsx",
          [("Product", "Unit", "Price"), ("Apples", "kg", 2.4), ("Bananas", "kg", 1.85), ("Cherries", "500 g", 6.9),
           ("Dates", "250 g", 4.15), ("Elderberries", "250 g", 7.3)])


@oracle("sheets-edit-existing")
def _(ctx: Ctx):
    import openpyxl
    book = openpyxl.load_workbook(ctx.p("budget.xlsx"))
    book.active.append(("Travel", 1200))
    book.active.append(("Total", "=SUM(B2:B5)"))
    book.save(ctx.p("budget.xlsx"))


@oracle("sheets-birthdays")
def _(ctx: Ctx):
    from .t_sheets import TEAM
    _xlsx(ctx.p("birthdays.xlsx"), [("Name", "Birthday")] + list(TEAM.items()))


# --- calendar, reminders, notes, mail ------------------------------------------------------------------------

@oracle("pim-reminder")
def _(ctx: Ctx):
    ap.reminders_add(ctx.apple, "MintBench: renew passport", tomorrow_at(10))


@oracle("pim-reminder-followup")
def _(ctx: Ctx):
    ap.reminders_add(ctx.apple, "MintBench: call the plumber", tomorrow_at(11), "bring the invoice")


@oracle("pim-notes-to-reminders")
def _(ctx: Ctx):
    for line in ("Buy printer ink", "Email the landlord", "Book the dentist"):
        ap.reminders_add(ctx.apple, line)


@oracle("pim-calendar-event")
def _(ctx: Ctx):
    start = next_weekday(1, 15)
    ap.calendar_add(ctx.apple, "MintBench sync", start, start.replace(hour=16))


@oracle("pim-calendar-conflict")
def _(ctx: Ctx):
    ap.calendar_add(ctx.apple, "MintBench 1:1 with Sam", tomorrow_at(15), tomorrow_at(15, 30))


@oracle("pim-note")
def _(ctx: Ctx):
    ap.notes_add(ctx.apple, "MintBench groceries", ["- eggs", "- milk", "- rice"])


def _draft(ctx: Ctx, subject: str, body: str) -> None:
    """An unsent, unsaved, invisible compose message, like the draft Mint should leave."""
    ap.mail_make_draft(ctx.apple, subject, "bench@example.com", body)


@oracle("pim-email-draft")
def _(ctx: Ctx):
    _draft(ctx, "[MintBench] Weekly status", "Hi, this week the website redesign went live, the Berlin lease is "
                                             "signed, and the data migration slipped a week.")


# --- web ------------------------------------------------------------------------------------------------------

@oracle("web-local-fact")
def _(ctx: Ctx):
    sb.write("answer.txt", "37 days\n")


@oracle("web-compare")
def _(ctx: Ctx):
    sb.write("laptops.md", "| Laptop | Price | Weight | Battery |\n|---|---|---|---|\n"
             "| Beta Book 14 | $999 | 1.42 kg | 14 h |\n| Alpha Air 13 | $1,149 | 1.19 kg | 17 h |\n"
             "| Gamma Pro 16 | $1,899 | 2.05 kg | 21 h |\n")


@oracle("web-slow-page")
def _(ctx: Ctx):
    sb.write("q3.txt", "Revenue: $4.82M\nNew customers: 1,375\nChurn: 2.9%\n")


@oracle("web-flaky-page")
def _(ctx: Ctx):
    for _attempt in range(3):
        try:
            urllib.request.urlopen(ctx.site + "/flaky/status.html", timeout=10).read()
            break
        except OSError:
            continue
    sb.write("version.txt", "4.7.2\n")


@oracle("web-public-fact")
def _(ctx: Ctx):
    sb.write("python.txt", "Python 3.0: 2008\n")


@oracle("web-follow-links")
def _(ctx: Ctx):
    sb.write("wiki.txt", "Port Veyra, 1742\n")


# --- clipboard --------------------------------------------------------------------------------------------------

@oracle("clip-copy-path")
def _(ctx: Ctx):
    sb.clipboard_set(str(ctx.p("Docs", "contract.pdf")))


@oracle("clip-pin")
def _(ctx: Ctx):
    import json
    from .t_misc import PIN_TEXT
    pins = sb.pins()
    pins["MintBench template"] = {"kind": "text", "text": PIN_TEXT}
    sb.PINS.parent.mkdir(parents=True, exist_ok=True)
    sb.PINS.write_text(json.dumps(pins))


@oracle("clip-history")
def _(ctx: Ctx):
    sb.clipboard_set("MintBench code ALPHA-1111")


@oracle("clip-to-file")
def _(ctx: Ctx):
    sb.write("scores.csv", sb.clipboard_get())


# --- video ----------------------------------------------------------------------------------------------------------

def _ffmpeg(*args: str) -> None:
    subprocess.run(["ffmpeg", "-v", "error", "-y", *args], check=True, capture_output=True)


@oracle("video-trim")
def _(ctx: Ctx):
    _ffmpeg("-ss", "5", "-i", str(ctx.p("Video", "demo.mp4")), "-c:v", "libx264", "-c:a", "aac",
            str(ctx.p("Video", "demo (edited).mp4")))


@oracle("video-extract-audio")
def _(ctx: Ctx):
    _ffmpeg("-i", str(ctx.p("Video", "demo.mp4")), "-vn", str(ctx.p("Video", "demo.mp3")))


@oracle("video-vertical")
def _(ctx: Ctx):
    _ffmpeg("-i", str(ctx.p("Video", "demo.mp4")), "-vf", "crop=ih*9/16:ih,scale=720:1280", "-c:a", "copy",
            str(ctx.p("Video", "demo (edited).mp4")))


@oracle("video-info-file")
def _(ctx: Ctx):
    sb.write("Video/video-info.txt", "clip.mov: 1280x720, 12 seconds\n")


# --- memory --------------------------------------------------------------------------------------------------------

@oracle("mem-remember-recall")
def _(ctx: Ctx):
    sb.bank_add("The user's MintBench project staging server is called kestrel-7.")
    sb.write("server.txt", "kestrel-7\n")


@oracle("mem-update")
def _(ctx: Ctx):
    sb.bank_remove(("9:30",))
    sb.bank_add("The user's MintBench team standup is at 10:15 every weekday.")


@oracle("mem-forget")
def _(ctx: Ctx):
    sb.bank_remove(("durian",))


@oracle("mem-preference-applied")
def _(ctx: Ctx):
    sb.bank_add("For MintBench work the user wants files as Markdown, saved in ~/MintBench/Notes.")
    sb.write("Notes/colours.md", "# Favourite colours\n\n- teal\n- amber\n- plum\n")


# --- error recovery ----------------------------------------------------------------------------------------------------

@oracle("err-write-conflict")
def _(ctx: Ctx):
    with open(ctx.p("plan.md"), "a") as f:
        f.write("\n## Checklist\n- [ ] Book the venue\n- [ ] Send the invites\n- [ ] Order the cake\n")


@oracle("err-interrupt-resume")
def _(ctx: Ctx):
    for i, day in enumerate(["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"], 1):
        sb.write(f"Week/day{i}.txt", day + "\n")


@oracle("err-path-typo")
def _(ctx: Ctx):
    sb.write("total.txt", "Total: 18,450\n")


@oracle("err-wrong-format")
def _(ctx: Ctx):
    _xlsx(ctx.p("Data", "scores.xlsx"), [("student", "score"), ("Ana", 91), ("Bo", 78), ("Cy", 85), ("Dee", 66)])


@oracle("err-rename-collision")
def _(ctx: Ctx):
    ctx.p("Drafts", "final.txt").rename(ctx.p("Drafts", "final-old.txt"))
    ctx.p("Drafts", "draft.txt").rename(ctx.p("Drafts", "final.txt"))


# --- multi-app ------------------------------------------------------------------------------------------------------------

@oracle("multi-research-report")
def _(ctx: Ctx):
    _xlsx(ctx.p("Research", "laptops.xlsx"), [("Laptop", "Price", "Weight", "Battery"), ("Alpha Air 13", 1149, 1.19, 17),
                                              ("Beta Book 14", 999, 1.42, 14), ("Gamma Pro 16", 1899, 2.05, 21)])
    sb.write("Research/recommendation.md", "Pick the Gamma Pro 16: 21 hours of battery.\n")
    ap.notes_add(ctx.apple, "MintBench laptop pick", ["Gamma Pro 16"])


@oracle("multi-invoice-pipeline")
def _(ctx: Ctx):
    _xlsx(ctx.p("invoices.xlsx"), [("Number", "Client", "Amount"), ("INV-101", "Acme Corp", 450),
                                   ("INV-102", "Globex", 1280), ("INV-103", "Initech", 2050.5),
                                   ("INV-104", "Umbrella", 999.99)])
    for n in ("INV-102", "INV-103"):
        _move(ctx.p("Invoices", f"{n}.txt"), ctx.p("Invoices", "Large", f"{n}.txt"))
    _draft(ctx, "[MintBench] Large invoices", "Large invoices: INV-102 (Globex, $1,280.00) and INV-103 (Initech, "
                                              "$2,050.50).")


@oracle("multi-screenshot-ocr")
def _(ctx: Ctx):
    from .t_files import image_bytes
    sb.write("mintbench-poster.png", image_bytes("white"))
    sb.write("poster.txt", "Zebra Quartz 7781\nDoors open at 19:45 - Hall C\nBring this code to the desk: KX-4419\n")


@oracle("multi-browser-form")
def _(ctx: Ctx):
    urllib.request.urlopen(ctx.site + "/record?name=MintBench+Tester&email=tester%40example.com&colour=Blue"
                           "&notes=&submitted=", timeout=10).read()


@oracle("multi-agent-astra")
def _(ctx: Ctx):
    sb.write("Agents/local-first.md", "# Local-first software\n\n" + " ".join(
        ["Local-first software keeps your data on your own device first, syncing when it can."] * 12) + "\n")



# --- harder tasks (rb/t_hard2.py) ------------------------------------------------------------------------------------------

@oracle("hard-shop-cheapest")
def _(ctx: Ctx):
    _xlsx(ctx.p("Shop", "cheapest.xlsx"), [("Product", "Price"), ("Seed Tray", 3.15), ("Bamboo Spatula", 3.49),
                                           ("Sticky Notes", 3.75), ("Total", "=SUM(B2:B4)")])
    sb.write("Shop/summary.md", "Cheapest in stock: Seed Tray at $3.15.\nThe three together cost $10.39.\n")


@oracle("hard-branch-404")
def _(ctx: Ctx):
    sb.write("hours.txt", "Riverside Quay\nMonday to Friday: 8:30 am - 6 pm\nSaturday: 10:00 - 16:00\nSunday: closed\n")


@oracle("hard-sales-clean")
def _(ctx: Ctx):
    from .t_hard2 import SALES_CLEAN, SALES_TOTAL
    sb.write("Data/sales-clean.csv", "order,date,customer,amount\n" + "".join(",".join(r) + "\n" for r in SALES_CLEAN))
    sb.write("Data/total.txt", f"Total: ${SALES_TOTAL}\n")


@oracle("hard-recon-pdf-csv")
def _(ctx: Ctx):
    sb.write("Recon/mismatches.csv", "invoice,statement,ledger,difference\nNW-2202,1275.40,1257.40,18.00\n"
             "NW-2205,460.25,406.25,54.00\nNW-2206,118.00,,\n")


@oracle("hard-acme-merge")
def _(ctx: Ctx):
    clients = ctx.p("Clients")
    for src, dst in (("acme_old/contract.pdf", "Acme/older/contract.pdf"), ("acme_old/brief.txt", "Acme/older/brief.txt"),
                     ("Acme Corp (copy)/brief.txt", "Acme/brief.txt"),
                     ("Acme Corp (copy)/invoice-0917.pdf", "Acme/invoice-0917.pdf"),
                     ("acme_old/2025/notes.txt", "Acme/notes.txt")):
        _move(clients / src, clients / dst)
    for gone in ("acme_old", "Acme Corp (copy)"):
        shutil.rmtree(clients / gone)


@oracle("hard-week-digest")
def _(ctx: Ctx):
    from .t_hard2 import LOGS_IN
    parts = [f"## {name}\n\n{ctx.p('Logs', name).read_text()}" for name in LOGS_IN]
    sb.write("Logs/this-week.md", "# This week\n\n" + "\n".join(parts))


@oracle("hard-spanish-brief")
def _(ctx: Ctx):
    sb.write("Docs/resumen.md", "# Lumen 4.12.0 (2026-09-21)\n\n"
             "- Modo sin conexión en la app móvil: las listas y los informes se abren sin red y se sincronizan después.\n"
             "- Exportación a CSV de todos los informes desde el menú Compartir.\n"
             "- La búsqueda es un 30 % más rápida.\n"
             "- Se corrigen el bucle de inicio de sesión en Safari y la hora de los recordatorios.\n\n"
             "La API v1 se apagará el 2026-12-01: hay que pasar a la API v2 antes de esa fecha.\n")


@oracle("hard-trip-euros")
def _(ctx: Ctx):
    _xlsx(ctx.p("Trip", "lisbon.xlsx"), [("Item", "Amount (EUR)"), ("Flights", 386.4), ("Hotel (3 nights)", 262.2),
                                         ("Food (4 days)", 147.2), ("Museum pass", 26.68), ("Airport taxi", 32.2),
                                         ("Total", "=SUM(B2:B6)")])


@oracle("hard-second-draft")
def _(ctx: Ctx):
    sb.write("Pitches/pitch-short.md", "# Pitch\n\nOrbitra is the simplest way to run your team's week. It costs "
             "just $49 a month. Setup takes minutes.\n")


@oracle("hard-invoice-dated")
def _(ctx: Ctx):
    from .t_hard2 import SEPT
    for name in SEPT:
        _move(ctx.p("Invoices", name), ctx.p("Invoices", "2026-09", name))


@oracle("hard-undo-replace")
def _(ctx: Ctx):
    path = ctx.p("Plan", "q-plan.md")
    path.write_text(path.read_text().replace("# Q3 plan", "# Q4 plan", 1))


@oracle("hard-undo-move-copy")
def _(ctx: Ctx):
    shutil.copy2(ctx.p("Share", "report-final.pdf"), ctx.p("Archive", "report-final-2026-09.pdf"))


@oracle("hard-specs-locked")
def _(ctx: Ctx):
    from .t_hard2 import STAMP
    for name in ("spec-a.md", "spec-b.md", "spec-d.md"):
        with open(ctx.p("Specs", name), "a") as f:
            f.write(f"\n{STAMP}\n")
    sb.write("Specs/skipped.txt", "spec-c.md (locked)\nspec-e.md (read-only)\n")


@oracle("hard-astra-grinders")
def _(ctx: Ctx):
    sb.write("Research/grinders.md", "# Coffee grinders\n\n| Grinder | Price | Burrs | Grind settings |\n|---|---|---|---|\n"
             "| Burrline S2 | $189 | Stainless steel, conical | 40 |\n| Ember Mill | $129 | Ceramic, flat | 18 |\n"
             "| Tidewell Pro | $159 | Hardened steel, flat | 55 |\n\nPick: Tidewell Pro\n")


@oracle("hard-clip-three")
def _(ctx: Ctx):
    from .t_hard2 import CLIPS
    with open(ctx.p("notes.md"), "a") as f:
        f.write("".join(f"- {c}\n" for c in CLIPS[1:]))


@oracle("hard-inventory-edit")
def _(ctx: Ctx):
    import openpyxl
    path = ctx.p("Shop", "inventory.xlsx")
    book = openpyxl.load_workbook(path)
    sheet = book["Stock"]
    sheet["B2"] = 14
    sheet.append(("Green Mug", 20, 8.5))
    sheet.append(("Total", "=SUM(B2:B6)"))
    book.save(path)


@oracle("hard-video-chain")
def _(ctx: Ctx):
    video = ctx.p("Video")
    _ffmpeg("-ss", "5", "-t", "10", "-i", str(video / "talk.mp4"), "-c:v", "libx264", "-c:a", "aac",
            str(video / "clip.mp4"))
    _ffmpeg("-i", str(video / "clip.mp4"), "-vn", str(video / "clip.mp3"))
    sb.write("Video/durations.txt", "clip.mp4: 10.0 seconds\nclip.mp3: 0:10\n")


@oracle("hard-accountant-memory")
def _(ctx: Ctx):
    sb.bank_add("The user's MintBench accountant is Farida Okonkwo (farida@example.com).")
    sb.write("Receipts/for-accountant.md", "Hi Farida,\n\nThere are 4 receipts this month, totalling $209.45.\n")


@oracle("hard-reminders-csv")
def _(ctx: Ctx):
    import datetime as dt
    from .t_hard2 import REMINDERS_WANT
    titles = {"domain": "MintBench: renew the domain", "deck": "MintBench: send the Q3 deck",
              "toner": "MintBench: order toner"}
    for key, days in REMINDERS_WANT.items():
        ap.reminders_add(ctx.apple, titles[key], dt.datetime.combine(dt.date.today() + dt.timedelta(days=days),
                                                                     dt.time(9)))


@oracle("hard-rates-convert")
def _(ctx: Ctx):
    from .t_hard2 import GBP, PRICES_USD
    sb.write("Shop/prices-gbp.csv", "sku,product,price_gbp\n" +
             "".join(f"{s},{p},{round(v * GBP + 1e-9, 2):.2f}\n" for s, p, v in PRICES_USD))
