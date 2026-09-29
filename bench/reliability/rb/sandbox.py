"""File-system sandbox (~/MintBench), stray outputs, the clipboard, Mint's memory and pins, and
readers for the files Mint makes (xlsx, docx, pdf, rtf, media)."""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import re
import shutil
import subprocess
import time
import zipfile
from pathlib import Path

from . import mintlink

HOME = Path.home()
ROOT = Path(os.environ.get("MINTBENCH_ROOT", HOME / "MintBench"))
MARKERS = ("mintbench", "mint bench")


def storage_folder() -> Path:
    """Mint's own Documents folder (Settings > Storage; default ~/Documents/Mint)."""
    try:
        chosen = json.loads((mintlink.RUNTIME / "settings.json").read_text()).get("storage_folder") or ""
    except (OSError, ValueError):
        chosen = ""
    return Path(chosen).expanduser() if str(chosen).strip() else HOME / "Documents" / "Mint"


# --- the sandbox folder -----------------------------------------------------------------------

def reset_root() -> None:
    wipe_root()
    ROOT.mkdir(parents=True, exist_ok=True)


def wipe_root() -> None:
    if not ROOT.exists():
        return
    # Read-only files and locked (uchg) files from the fixtures must not stop the wipe.
    subprocess.run(["chflags", "-R", "nouchg", str(ROOT)], capture_output=True, check=False)
    subprocess.run(["chmod", "-R", "u+rwX", str(ROOT)], capture_output=True, check=False)
    shutil.rmtree(ROOT, ignore_errors=True)


def write(rel: str, content: str | bytes, mtime: dt.datetime | None = None) -> Path:
    path = ROOT / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(content, bytes):
        path.write_bytes(content)
    else:
        path.write_text(content)
    if mtime is not None:
        stamp = mtime.timestamp()
        os.utime(path, (stamp, stamp))
    return path


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def files_under(folder: Path) -> list[Path]:
    if not folder.exists():
        return []
    return sorted(p for p in folder.rglob("*") if p.is_file() and not p.name.startswith("."))


def born_after(path: Path, since: float) -> bool:
    st = path.stat()
    return max(getattr(st, "st_birthtime", st.st_mtime), st.st_mtime) >= since - 1


def new_files(since: float, suffixes: tuple[str, ...] = (), where: tuple[Path, ...] = ()) -> list[Path]:
    """Files made or changed since `since` in ~/MintBench and Mint's Documents folder (or `where`)."""
    found = []
    for folder in where or (ROOT, storage_folder()):
        for path in files_under(folder):
            if suffixes and path.suffix.lower() not in suffixes:
                continue
            try:
                if born_after(path, since):
                    found.append(path)
            except OSError:
                continue
    return found


# --- outputs Mint put outside the sandbox ------------------------------------------------------

def stray_places() -> list[tuple[Path, bool]]:
    """(folder, recursive) where outputs that should have gone into ~/MintBench tend to land."""
    return [(storage_folder(), True), (HOME / "Desktop", False), (HOME / "Downloads", False),
            (HOME / "Documents", False), (HOME / "Movies", False), (HOME / "Pictures", False)]


def find_strays(since: float, markers: tuple[str, ...]) -> list[Path]:
    """Files made since `since` outside the sandbox whose NAME contains one of `markers`."""
    marks = tuple(m.lower() for m in (*MARKERS, *markers) if m)
    found = []
    for folder, recursive in stray_places():
        if not folder.exists():
            continue
        paths = folder.rglob("*") if recursive else folder.iterdir()
        for path in paths:
            try:
                if ROOT in path.parents or not path.is_file():
                    continue
                if any(m in path.name.lower() for m in marks) and born_after(path, since):
                    found.append(path)
            except OSError:
                continue
    return found


def remove_strays(paths: list[Path]) -> list[str]:
    removed = []
    for path in paths:
        try:
            path.unlink()
            removed.append(str(path).replace(str(HOME), "~"))
            parent = path.parent
            if parent != storage_folder() and parent.parent != HOME and not any(parent.iterdir()):
                parent.rmdir()
        except OSError:
            pass
    return removed


def purge_trash(fixtures: dict[str, str]) -> list[str]:
    """Remove from ~/.Trash the fixture files Mint trashed: name -> sha256. Only exact content matches
    are removed, and only by known name (the Trash itself cannot be listed without Full Disk Access)."""
    trash = HOME / ".Trash"
    removed = []
    for name, digest in fixtures.items():
        stem, suffix = os.path.splitext(name)
        candidates = [trash / name] + [trash / f"{stem} {n}{suffix}" for n in range(2, 6)]
        for path in candidates:
            try:
                if path.is_file() and sha(path) == digest:
                    path.unlink()
                    removed.append(path.name)
            except OSError:
                continue
    return removed


# --- the clipboard ---------------------------------------------------------------------------------

def clipboard_get() -> str:
    import AppKit
    return str(AppKit.NSPasteboard.generalPasteboard().stringForType_(AppKit.NSPasteboardTypeString) or "")


BENCH_CLIPS: set[str] = set()      # every text the benchmark put on the clipboard


def clipboard_is_bench() -> bool:
    """Does the clipboard hold the benchmark's text (so restoring the user's is safe)? Anything else may be
    something the user copied during the run, and is left alone."""
    text = clipboard_get()
    files = clipboard_files()
    if files:
        return all(str(ROOT) in f for f in files)
    return bool(text) and (text in BENCH_CLIPS or "mintbench" in text.lower() or str(ROOT) in text)


def clipboard_set(text: str) -> None:
    BENCH_CLIPS.add(text)
    import AppKit
    board = AppKit.NSPasteboard.generalPasteboard()
    board.clearContents()
    board.setString_forType_(text, AppKit.NSPasteboardTypeString)


def clipboard_files() -> list[str]:
    import AppKit
    board = AppKit.NSPasteboard.generalPasteboard()
    urls = board.readObjectsForClasses_options_([AppKit.NSURL], {AppKit.NSPasteboardURLReadingFileURLsOnlyKey: True})
    return [str(u.path()) for u in (urls or [])]


# --- Mint's memory bank and clipboard pins (files re-read by Mint on every use) --------------------

BANK = mintlink.RUNTIME / "memory" / "bank.json"
VIEW = mintlink.RUNTIME / "memory" / "MEMORY.md"
PINS = mintlink.RUNTIME / "clipboard" / "pins.json"


def _atomic_json(path: Path, data) -> None:
    tmp = path.with_suffix(".mintbench.tmp")
    tmp.write_text(json.dumps(data, indent=1, ensure_ascii=False))
    os.chmod(tmp, 0o600)
    tmp.replace(path)


def bank() -> list[dict]:
    try:
        data = json.loads(BANK.read_text())
        return data if isinstance(data, list) else []
    except (OSError, ValueError):
        return []


def bank_add(text: str, group: str = "work") -> str:
    blocks = bank()
    n = max([int(b["id"][1:]) for b in blocks if re.match(r"^m\d+$", str(b.get("id")))] + [0]) + 1
    now = time.strftime("%Y-%m-%d %H:%M")
    blocks.append({"id": f"m{n}", "group": group, "text": text, "pinned": False, "created": now,
                   "updated": now, "source": "user", "hits": 0})
    _atomic_json(BANK, blocks)
    return f"m{n}"


def bank_matching(words: tuple[str, ...]) -> list[dict]:
    words = tuple(w.lower() for w in words)
    return [b for b in bank() if any(w in str(b.get("text", "")).lower() for w in words)]


def bank_remove(words: tuple[str, ...]) -> list[str]:
    """Remove every memory whose text has one of `words`; also from the MEMORY.md view."""
    words = tuple(w.lower() for w in words if w)
    blocks = bank()
    keep = [b for b in blocks if not any(w in str(b.get("text", "")).lower() for w in words)]
    gone = [str(b.get("text", ""))[:120] for b in blocks if b not in keep]
    if gone:
        _atomic_json(BANK, keep)
        try:
            lines = VIEW.read_text().splitlines()
            VIEW.write_text("\n".join(x for x in lines if not any(w in x.lower() for w in words)) + "\n")
        except OSError:
            pass
    return gone


def pins() -> dict:
    try:
        data = json.loads(PINS.read_text())
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def pins_remove(words: tuple[str, ...]) -> list[str]:
    words = tuple(w.lower() for w in words if w)
    current = pins()
    gone = [k for k, v in current.items() if any(w in (k + " " + json.dumps(v)).lower() for w in words)]
    if gone:
        _atomic_json(PINS, {k: v for k, v in current.items() if k not in gone})
    return gone


def skills_snapshot() -> set[str]:
    folder = mintlink.RUNTIME / "skills"
    return {str(p.relative_to(folder)) for p in folder.rglob("*.md")} if folder.exists() else set()


# --- readers ------------------------------------------------------------------------------------------

def xlsx_rows(path: Path) -> list[list]:
    """Every row of every sheet (values, not formulas where a cached value exists)."""
    import openpyxl
    rows = []
    for values_only in (True,):
        book = openpyxl.load_workbook(path, data_only=values_only, read_only=True)
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() for c in row):
                    rows.append(list(row))
        book.close()
    return rows


def xlsx_formulas(path: Path) -> list[list]:
    import openpyxl
    book = openpyxl.load_workbook(path, data_only=False, read_only=True)
    rows = [list(r) for s in book.worksheets for r in s.iter_rows(values_only=True)]
    book.close()
    return rows


def docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    xml = re.sub(r"</w:p>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml)


def pdf_text(path: Path) -> str:
    from Foundation import NSURL
    from Quartz import PDFDocument
    doc = PDFDocument.alloc().initWithURL_(NSURL.fileURLWithPath_(str(path)))
    return str(doc.string() or "") if doc is not None else ""


def rich_text(path: Path) -> str:
    """Text of rtf/doc/docx/html/pages-exported files through textutil."""
    done = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)], capture_output=True, text=True)
    return done.stdout


def any_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return pdf_text(path)
        if suffix == ".docx":
            return docx_text(path)
        if suffix in (".rtf", ".rtfd", ".doc", ".html", ".htm", ".webarchive", ".odt"):
            return rich_text(path)
        if suffix == ".xlsx":
            return "\n".join(" | ".join("" if c is None else str(c) for c in r) for r in xlsx_rows(path))
        return path.read_text(errors="replace")
    except Exception as error:  # noqa: BLE001 - a broken output is a failed check, not a crash
        return f"<unreadable: {error}>"


def media(path: Path) -> dict:
    """ffprobe: duration (s), width, height, has_audio, has_video."""
    done = subprocess.run(["ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams",
                           str(path)], capture_output=True, text=True)
    try:
        info = json.loads(done.stdout)
    except ValueError:
        return {}
    video = next((s for s in info.get("streams", []) if s.get("codec_type") == "video"), None)
    audio = next((s for s in info.get("streams", []) if s.get("codec_type") == "audio"), None)
    return {"duration": float(info.get("format", {}).get("duration") or 0),
            "width": int(video.get("width", 0)) if video else 0,
            "height": int(video.get("height", 0)) if video else 0,
            "has_video": video is not None, "has_audio": audio is not None}


def make_video(path: Path, seconds: float, size: str = "1280x720", tone: int = 440) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", f"testsrc2=size={size}:rate=30",
                    "-f", "lavfi", "-i", f"sine=frequency={tone}:sample_rate=44100", "-t", str(seconds),
                    "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", str(path)],
                   check=True, capture_output=True)
    return path


def launched_apps() -> set[str]:
    import AppKit
    return {str(a.localizedName()) for a in AppKit.NSWorkspace.sharedWorkspace().runningApplications()
            if a.activationPolicy() == AppKit.NSApplicationActivationPolicyRegular}


# Document viewers and fixture apps a task may open. Only these are ever asked to quit, and only if the task started them.
VIEWERS = {"TextEdit", "Preview", "Numbers", "Pages", "Microsoft Excel", "Microsoft Word", "QuickTime Player",
           "LibreOffice", "Keynote", "Reminders", "Calendar", "Notes", "Mail"}


def quit_apps(names: set[str]) -> list[str]:
    """Politely quit (no force: unsaved work gets its dialog) the given apps."""
    import AppKit
    asked = []
    for app in AppKit.NSWorkspace.sharedWorkspace().runningApplications():
        if str(app.localizedName()) in names:
            app.terminate()
            asked.append(str(app.localizedName()))
    return asked
